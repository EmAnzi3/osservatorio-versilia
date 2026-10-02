#!/usr/bin/env python3
from __future__ import annotations
import argparse, io, json, zipfile
from pathlib import Path
import requests
from openpyxl import load_workbook

URL="https://www.istat.it/wp-content/uploads/2026/01/Tavole.zip"

def find_row(sheet,label):
    for r in range(1,(sheet.max_row or 0)+1):
        v=sheet.cell(r,1).value
        if str(v or "").strip().lower()==label.lower():
            return r
    return None

def main():
    ap=argparse.ArgumentParser(); ap.add_argument("--output",required=True); a=ap.parse_args()
    s=requests.Session(); s.headers["User-Agent"]="OsservatorioVersilia-A3-benchmark-acquire/1.0"
    r=s.get(URL,timeout=240); r.raise_for_status()
    z=zipfile.ZipFile(io.BytesIO(r.content))
    member="Tavole/1. Appendice Statistica Frame territoriale.xlsx"
    wb=load_workbook(io.BytesIO(z.read(member)),read_only=True,data_only=True)
    sheet=None
    for sh in wb.worksheets:
        vals=" ".join(str(sh.cell(rr,1).value or "") for rr in range(1,min(8,sh.max_row or 0)+1)).lower()
        if "principali aggregati" in vals and find_row(sh,"Toscana") and find_row(sh,"ITALIA"):
            sheet=sh; break
    if sheet is None: raise RuntimeError("Tavola regionale Frame-SBS non individuata")
    hdr=None
    for rr in range(1,12):
        row=[sheet.cell(rr,c).value for c in range(1,(sheet.max_column or 0)+1)]
        if any("Numero unità locali" in str(v or "") for v in row) and any(str(v or "").strip()=="Addetti" for v in row):
            hdr=row; break
    if hdr is None: raise RuntimeError("Header Frame-SBS non individuato")
    def col(name):
        for i,v in enumerate(hdr,1):
            if str(v or "").strip()==name: return i
        raise RuntimeError(f"Colonna mancante: {name}")
    cu,ca=col("Numero unità locali"),col("Addetti")
    out={}
    for label,key in (("Toscana","tuscany"),("ITALIA","italy")):
        rr=find_row(sheet,label)
        units=float(sheet.cell(rr,cu).value); employees=float(sheet.cell(rr,ca).value)
        out[key]={"localUnits":units,"localEmployees":employees,"employeesPerLocalUnit":employees/units}
    payload={
      "schemaVersion":2,"publisher":"Istat","profileId":"istat-business-annual","year":2023,
      "sourceUrl":URL,"archiveMember":member,"sheet":sheet.title,
      "status":"SOURCE_LINEAGE_MISMATCH",
      "blockedCandidates":out,
      "reason":(
        "Le metriche pubbliche localUnits/localEmployees/employeesPerLocalUnit usano ASIA-UL; "
        "la Tavola Frame-SBS regionale usa il perimetro Frame territoriale e la grandezza Addetti. "
        "Non si promuovono benchmark finché Toscana/Italia non sono acquisiti dallo stesso flusso ASIA-UL."
      ),
      "formulas":{"frameEmployeesPerLocalUnit":"Addetti / Numero unità locali"}
    }
    p=Path(a.output); p.parent.mkdir(parents=True,exist_ok=True); p.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"status":payload["status"],"blockedMetrics":["localUnits","localEmployees","employeesPerLocalUnit"],"output":str(p)},ensure_ascii=False))
if __name__=="__main__": main()
