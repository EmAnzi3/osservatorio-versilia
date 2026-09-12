#!/usr/bin/env python3
"""Genera lo snapshot ufficiale v1.37 da fonti ISPRA/Regione Toscana/Istat.

Il generatore e' destinato alla sola fase di acquisizione/QA. La build pubblica usa
esclusivamente il JSON versionato prodotto da questo script e non accede alla rete.
"""
from __future__ import annotations

import hashlib
import io
import json
import math
import shutil
import tempfile
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import geopandas as gpd
from openpyxl import load_workbook
from shapely.ops import unary_union

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "source-snapshots" / "territorio-v137-official.json"
SITE_DATA = ROOT / "data" / "site-data.json"
OLD_BONIFICA = ROOT / "data" / "source-snapshots" / "bonifica-rischio-v126-gis.json"

TOWNS = [
    "Camaiore", "Forte dei Marmi", "Massarosa", "Pietrasanta",
    "Seravezza", "Stazzema", "Viareggio",
]
TOWN_CODES = {
    "Camaiore": "046005", "Forte dei Marmi": "046013", "Massarosa": "046018",
    "Pietrasanta": "046024", "Seravezza": "046028", "Stazzema": "046030",
    "Viareggio": "046033",
}
SOIL_YEARS = [2006, 2012, 2015, 2016, 2017, 2018, 2019, 2020, 2021, 2022, 2023, 2024]
CHANGE_YEARS = [2012, 2015, 2016, 2017, 2018, 2019, 2020, 2021, 2022, 2023, 2024]
UA = {"User-Agent": "OsservatorioVersilia-v137-snapshot/1.0"}

ISPRA_FULL = "https://groupware.sinanet.isprambiente.it/uso-copertura-e-consumo-di-suolo/library/consumo-di-suolo/indicatori/consumo_suolo_2025_com_prov_reg_naz_v1.1/download/en/1/consumo_suolo_2025_Com_Prov_Reg_Naz_v1.1.xlsx"
ISTAT_BOUNDARIES = "https://www.istat.it/storage/cartografia/confini_amministrativi/generalizzati/2026/Limiti01012026_g.zip"
RETICULUM = "https://www.regione.toscana.it/documents/d/guest/infrastruttura_rev25-zip"
ROADS = "https://www502.regione.toscana.it/geoscopio/download/grafo_stradale/iternet.zip"
WFS_BASE = "https://www502.regione.toscana.it/wmsraster/com.rt.wms.RTmap/wms"
WFS_LAYERS = {
    "parksNational": "rt_arprot.idparnaz.rt.poly",
    "reservesNational": "rt_arprot.idrisnatstat.rt.poly",
    "parksRegional": "rt_arprot.idparreg.rt.poly",
    "parksProvincial": "rt_arprot.idparprov.rt.poly",
    "reservesRegional": "rt_arprot.idrisnatreg.rt.poly",
    "anpil": "rt_arprot.idanpil.rt.poly",
    "zsc": "rt_arprot.idnat2000_sic.rt.poly",
    "zps": "rt_arprot.idnat2000_zps.rt.poly",
    "zscZps": "rt_arprot.idnat2000_sic_zps.rt.poly",
    "ramsar": "rt_arprot.idramsar.rt.poly",
}


def download(url: str, dest: Path) -> dict:
    req = urllib.request.Request(url, headers=UA)
    hasher = hashlib.sha256()
    size = 0
    with urllib.request.urlopen(req, timeout=900) as response, dest.open("wb") as out:
        while True:
            chunk = response.read(1024 * 1024)
            if not chunk:
                break
            out.write(chunk)
            hasher.update(chunk)
            size += len(chunk)
    return {"url": url, "sizeBytes": size, "sha256": hasher.hexdigest()}


def wfs_url(layer: str) -> str:
    params = {
        "map": "wmsarprot", "service": "WFS", "version": "2.0.0",
        "request": "GetFeature", "OUTPUTFORMAT": "application/shapefile",
        "filename": "v137_snapshot", "typename": layer,
    }
    return WFS_BASE + "?" + urllib.parse.urlencode(params)


def safe_geom(geom):
    if geom is None or geom.is_empty:
        return geom
    if not geom.is_valid:
        try:
            geom = geom.make_valid()
        except AttributeError:
            geom = geom.buffer(0)
    return geom


def union_frame(frame: gpd.GeoDataFrame):
    geoms = [safe_geom(g) for g in frame.geometry if g is not None and not g.is_empty]
    geoms = [g for g in geoms if g is not None and not g.is_empty]
    return safe_geom(unary_union(geoms)) if geoms else None


def normalize_code(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    digits = "".join(ch for ch in text if ch.isdigit())
    return digits.zfill(6) if digits else ""


def load_boundaries(zip_path: Path, work: Path):
    folder = work / "boundaries"
    folder.mkdir()
    with zipfile.ZipFile(zip_path) as archive:
        archive.extractall(folder)
    candidates = list(folder.rglob("Com01012026_g_WGS84.shp")) or list(folder.rglob("Com*.shp"))
    if not candidates:
        raise RuntimeError("Istat: shapefile comunale non trovato")
    frame = gpd.read_file(candidates[0], engine="pyogrio")
    code_field = next((c for c in ("PRO_COM_T", "PRO_COM", "PRO_COM_T_", "COD_PROCOM") if c in frame.columns), None)
    if not code_field:
        raise RuntimeError(f"Istat: campo codice comunale non trovato: {list(frame.columns)}")
    frame = frame.copy()
    frame["_code"] = frame[code_field].map(normalize_code)
    wanted = set(TOWN_CODES.values())
    frame = frame[frame["_code"].isin(wanted)].copy()
    if set(frame["_code"]) != wanted:
        raise RuntimeError(f"Istat: confini 7/7 non trovati: {sorted(set(frame['_code']))}")
    frame = frame.to_crs(3003)
    geoms = {}
    for town in TOWNS:
        part = frame[frame["_code"] == TOWN_CODES[town]]
        geoms[town] = safe_geom(unary_union(list(part.geometry)))
    versilia = safe_geom(unary_union([geoms[t] for t in TOWNS]))
    areas_km2 = {town: geoms[town].area / 1_000_000 for town in TOWNS}
    return geoms, versilia, areas_km2, candidates[0].name


def read_population_series() -> dict[str, dict[int, float]]:
    site = json.loads(SITE_DATA.read_text(encoding="utf-8"))
    metric = site.get("metrics", {}).get("population")
    if not metric:
        raise RuntimeError("population: metrica canonica non trovata in data/site-data.json")
    result = {}
    for row in metric.get("rows", []):
        town = row.get("town")
        if town not in TOWNS:
            continue
        mapping = {}
        series = row.get("series")
        if isinstance(series, dict) and isinstance(series.get("years"), list) and isinstance(series.get("values"), list):
            mapping.update({int(y): float(v) for y, v in zip(series["years"], series["values"]) if v is not None})
        elif isinstance(series, list):
            for point in series:
                if isinstance(point, dict) and point.get("year") is not None and point.get("value") is not None:
                    mapping[int(point["year"])] = float(point["value"])
        current_year = metric.get("meta", {}).get("year")
        if row.get("value") is not None and current_year is not None:
            try:
                mapping[int(str(current_year)[:4])] = float(row["value"])
            except ValueError:
                pass
        result[town] = mapping
    if set(result) != set(TOWNS):
        raise RuntimeError("population: copertura comunale non 7/7")
    if not all(2024 in result[t] for t in TOWNS):
        raise RuntimeError("population: valore 2024 assente per almeno un Comune")
    return result


def workbook_rows(path: Path) -> dict[int, dict[str, dict]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    result = {}
    code_to_town = {int(code): town for town, code in TOWN_CODES.items()}
    for year in SOIL_YEARS:
        sheet_name = f"Comuni_{year}"
        if sheet_name not in wb.sheetnames:
            raise RuntimeError(f"ISPRA: foglio {sheet_name} assente")
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        headers = [str(x) if x is not None else "" for x in next(rows)]
        index = {name: i for i, name in enumerate(headers)}
        required = {"pro_com", "csuolo1", "csuolo2", "csuolo4"}
        if year in CHANGE_YEARS:
            required |= {"csuolo8", "csuolo10"}
        missing = required - set(index)
        if missing:
            raise RuntimeError(f"ISPRA {year}: campi assenti {sorted(missing)}")
        selected = {}
        for raw in rows:
            try:
                code_int = int(raw[index["pro_com"]])
            except (TypeError, ValueError):
                continue
            town = code_to_town.get(code_int)
            if town:
                selected[town] = {key: raw[pos] for key, pos in index.items()}
        if set(selected) != set(TOWNS):
            raise RuntimeError(f"ISPRA {year}: copertura {len(selected)}/7")
        result[year] = selected
    return result


def build_soil(ispra_path: Path, populations: dict[str, dict[int, float]]):
    data = workbook_rows(ispra_path)
    municipalities = {}
    changes = {}
    for town in TOWNS:
        stock_rows = {}
        change_rows = {}
        latest = data[2024][town]
        municipal_area_ha = float(latest["csuolo1"]) + float(latest["csuolo2"])
        for year in SOIL_YEARS:
            row = data[year][town]
            consumed = float(row["csuolo1"])
            pct = float(row["csuolo4"])
            pop = populations[town].get(year)
            sqm = consumed * 10_000 / pop if pop and pop > 0 else None
            stock_rows[str(year)] = {
                "consumedHa": round(consumed, 6),
                "consumedPct": round(pct, 6),
                "sqmPerResident": round(sqm, 6) if sqm is not None else None,
                "population": int(round(pop)) if pop is not None else None,
            }
            if year in CHANGE_YEARS:
                net = float(row["csuolo8"] or 0)
                gross = float(row["csuolo10"] or 0)
                change_rows[str(year)] = {"grossHa": round(gross, 6), "netHa": round(net, 6)}
        municipalities[town] = {"municipalAreaHa": round(municipal_area_ha, 6), "rows": stock_rows}
        changes[town] = {"rows": change_rows}

    versilia_rows = {}
    for year in SOIL_YEARS:
        consumed = sum(municipalities[t]["rows"][str(year)]["consumedHa"] for t in TOWNS)
        area = sum(float(data[year][t]["csuolo1"]) + float(data[year][t]["csuolo2"]) for t in TOWNS)
        pop_values = [populations[t].get(year) for t in TOWNS]
        total_pop = sum(pop_values) if all(v is not None for v in pop_values) else None
        versilia_rows[str(year)] = {
            "consumedHa": round(consumed, 6),
            "consumedPct": round(consumed / area * 100, 6),
            "sqmPerResident": round(consumed * 10_000 / total_pop, 6) if total_pop else None,
            "population": int(round(total_pop)) if total_pop else None,
        }
    change_versilia = {str(year): {
        "grossHa": round(sum(changes[t]["rows"][str(year)]["grossHa"] for t in TOWNS), 6),
        "netHa": round(sum(changes[t]["rows"][str(year)]["netHa"] for t in TOWNS), 6),
    } for year in CHANGE_YEARS}
    return (
        {"years": SOIL_YEARS, "municipalities": municipalities, "versilia": {"rows": versilia_rows}},
        {"years": CHANGE_YEARS, "municipalities": changes, "versilia": {"rows": change_versilia}},
    )


def extract_zip(path: Path, target: Path) -> None:
    target.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path) as archive:
        archive.extractall(target)


def load_reticulum(zip_path: Path, work: Path):
    folder = work / "reticulum"
    extract_zip(zip_path, folder)
    shp = next(iter(folder.rglob("reticoloDCR242025.shp")), None)
    if not shp:
        raise RuntimeError("Reticolo: reticoloDCR242025.shp non trovato")
    frame = gpd.read_file(shp, engine="pyogrio")
    if frame.crs is None:
        raise RuntimeError("Reticolo: CRS assente")
    frame = frame.to_crs(3003)
    if not {"COMPLR79", "RETGESLR79"}.issubset(frame.columns):
        raise RuntimeError("Reticolo: campi gestione assenti")
    managed = frame[(frame["COMPLR79"].astype(str) == "Toscana Nord") & (frame["RETGESLR79"].astype(str) == "SI")].copy()
    if len(frame) != 215923 or len(managed) != 26133:
        raise RuntimeError(f"Reticolo: conteggi sorgente inattesi total={len(frame)} managed={len(managed)}")
    return frame, managed, shp.name


def load_roads(outer_zip: Path, work: Path):
    outer = work / "roads-outer"
    extract_zip(outer_zip, outer)
    inner_zips = list(outer.rglob("*.zip"))
    if not inner_zips:
        raise RuntimeError("Iter.Net: archivio interno non trovato")
    inner = work / "roads-inner"
    extract_zip(inner_zips[0], inner)
    candidates = [p for p in inner.rglob("*.shp") if p.name.lower() == "elem_strad.shp"]
    if not candidates:
        raise RuntimeError("Iter.Net: elem_strad.shp non trovato")
    frame = gpd.read_file(candidates[0], engine="pyogrio", columns=[])
    if frame.crs is None:
        raise RuntimeError("Iter.Net: CRS assente")
    frame = frame.to_crs(3003)
    if len(frame) != 405341:
        raise RuntimeError(f"Iter.Net: feature count inatteso {len(frame)}")
    return frame, candidates[0].name


def clipped_line_km(frame: gpd.GeoDataFrame, polygon) -> float:
    minx, miny, maxx, maxy = polygon.bounds
    subset = frame.cx[minx:maxx, miny:maxy]
    if subset.empty:
        return 0.0
    lengths = subset.geometry.intersection(polygon).length
    return float(lengths.sum() / 1000)


def build_reticulum(full, managed, town_geoms, versilia_geom, areas_km2):
    municipalities = {}
    for town in TOWNS:
        full_km = clipped_line_km(full, town_geoms[town])
        managed_km = clipped_line_km(managed, town_geoms[town])
        municipalities[town] = {
            "municipalAreaKm2": round(areas_km2[town], 6),
            "fullNetworkKm": round(full_km, 6),
            "managedNetworkKm": round(managed_km, 6),
            "fullNetworkDensity": round(full_km / areas_km2[town], 6),
        }
    area = versilia_geom.area / 1_000_000
    full_km = clipped_line_km(full, versilia_geom)
    managed_km = clipped_line_km(managed, versilia_geom)
    return {
        "referenceLabel": "DCRT 24/2025 · confini Istat 1 gennaio 2026",
        "municipalities": municipalities,
        "versilia": {
            "unionAreaKm2": round(area, 6),
            "fullNetworkKm": round(full_km, 6),
            "managedNetworkKm": round(managed_km, 6),
            "fullNetworkDensity": round(full_km / area, 6),
        },
    }


def build_roads(frame, town_geoms, versilia_geom, areas_km2):
    municipalities = {}
    for town in TOWNS:
        km = clipped_line_km(frame, town_geoms[town])
        municipalities[town] = {
            "municipalAreaKm2": round(areas_km2[town], 6),
            "graphKm": round(km, 6),
            "graphDensity": round(km / areas_km2[town], 6),
        }
    area = versilia_geom.area / 1_000_000
    km = clipped_line_km(frame, versilia_geom)
    return {
        "referenceLabel": "Iter.Net 4.48 · giugno 2022",
        "municipalities": municipalities,
        "versilia": {"unionAreaKm2": round(area, 6), "graphKm": round(km, 6), "graphDensity": round(km / area, 6)},
    }


def load_protected(work: Path):
    frames = {}
    provenance = {}
    for key, layer in WFS_LAYERS.items():
        url = wfs_url(layer)
        req = urllib.request.Request(url, headers=UA)
        raw = urllib.request.urlopen(req, timeout=300).read()
        provenance[key] = {"typeName": layer, "sizeBytes": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}
        folder = work / f"protected-{key}"
        folder.mkdir()
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            archive.extractall(folder)
        shp = next(folder.glob("*.shp"))
        frame = gpd.read_file(shp, engine="pyogrio")
        if frame.crs is None:
            raise RuntimeError(f"Aree protette {key}: CRS assente")
        frames[key] = frame.to_crs(3003)
    return frames, provenance


def geom_union(*geoms):
    flat = [g for g in geoms if g is not None and not g.is_empty]
    return safe_geom(unary_union(flat)) if flat else None


def build_protected(frames, town_geoms, versilia_geom):
    unions = {key: union_frame(frame) for key, frame in frames.items()}
    parks = geom_union(unions["parksNational"], unions["reservesNational"], unions["parksRegional"], unions["parksProvincial"], unions["reservesRegional"])
    zsc = geom_union(unions["zsc"], unions["zscZps"])
    zps = geom_union(unions["zps"], unions["zscZps"])
    natura = geom_union(unions["zsc"], unions["zps"], unions["zscZps"])
    categories = {
        "parksReserves": parks,
        "anpil": unions["anpil"],
        "natura2000": natura,
        "zsc": zsc,
        "zps": zps,
        "ramsar": unions["ramsar"],
    }
    total = geom_union(parks, unions["anpil"], natura, unions["ramsar"])
    labels = {
        "parksReserves": "Parchi e riserve",
        "anpil": "ANPIL (regime transitorio)",
        "natura2000": "Rete Natura 2000",
        "zsc": "ZSC",
        "zps": "ZPS",
        "ramsar": "Zone Ramsar",
    }
    def area_ha(geom, mask):
        if geom is None or geom.is_empty:
            return 0.0
        return geom.intersection(mask).area / 10_000
    municipalities = {}
    for town in TOWNS:
        mask = town_geoms[town]
        municipalities[town] = {
            "municipalAreaHa": round(mask.area / 10_000, 6),
            "totalUnionHa": round(area_ha(total, mask), 6),
            "categoriesHa": {key: round(area_ha(geom, mask), 6) for key, geom in categories.items()},
        }
    return {
        "referenceLabel": "archivi geografici ufficiali Regione Toscana · consultati 12 settembre 2026",
        "categoryOrder": list(categories),
        "categoryLabels": labels,
        "municipalities": municipalities,
        "versilia": {
            "municipalAreaHa": round(versilia_geom.area / 10_000, 6),
            "totalUnionHa": round(area_ha(total, versilia_geom), 6),
            "categoriesHa": {key: round(area_ha(geom, versilia_geom), 6) for key, geom in categories.items()},
        },
        "definitionNote": "Il totale e' l'unione geometrica di parchi/riserve, ANPIL, Natura 2000 e Ramsar; le sovrapposizioni sono conteggiate una sola volta. Le ANPIL sono mantenute come categoria distinta in regime transitorio.",
    }


def validate_managed_against_existing(reticulum: dict) -> None:
    if not OLD_BONIFICA.exists():
        return
    old = json.loads(OLD_BONIFICA.read_text(encoding="utf-8"))["managedReticulum"]["byTown"]
    for town in TOWNS:
        current = reticulum["municipalities"][town]["managedNetworkKm"]
        expected = float(old[town]["km"])
        if not math.isclose(current, expected, abs_tol=0.003, rel_tol=0):
            raise RuntimeError(f"Reticolo gestito {town}: regressione {current} != {expected}")


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="territorio-v137-") as tmp:
        work = Path(tmp)
        files = {}
        paths = {
            "ispra": work / "ispra.xlsx",
            "boundaries": work / "istat.zip",
            "reticulum": work / "reticulum.zip",
            "roads": work / "roads.zip",
        }
        files["ispra"] = download(ISPRA_FULL, paths["ispra"])
        files["boundaries"] = download(ISTAT_BOUNDARIES, paths["boundaries"])
        files["reticulum"] = download(RETICULUM, paths["reticulum"])
        files["roads"] = download(ROADS, paths["roads"])

        town_geoms, versilia_geom, areas_km2, boundary_layer = load_boundaries(paths["boundaries"], work)
        populations = read_population_series()
        land_use, land_change = build_soil(paths["ispra"], populations)

        full_reticulum, managed_reticulum, reticulum_layer = load_reticulum(paths["reticulum"], work)
        reticulum = build_reticulum(full_reticulum, managed_reticulum, town_geoms, versilia_geom, areas_km2)
        validate_managed_against_existing(reticulum)

        road_frame, road_layer = load_roads(paths["roads"], work)
        roads = build_roads(road_frame, town_geoms, versilia_geom, areas_km2)

        protected_frames, protected_provenance = load_protected(work)
        protected = build_protected(protected_frames, town_geoms, versilia_geom)

        payload = {
            "schemaVersion": 1,
            "release": "v1.37.0",
            "verifiedAt": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
            "municipalityOrder": TOWNS,
            "sources": {
                "ispraSoil": {
                    "publisher": "ISPRA",
                    "page": "https://www.isprambiente.gov.it/it/attivita/suolo-e-territorio/suolo/il-consumo-di-suolo/i-dati-sul-consumo-di-suolo",
                    "reference": "Consumo di suolo, dinamiche territoriali e servizi ecosistemici · edizione 2025 · dati comunali 2006–2024",
                    "acquisition": files["ispra"],
                },
                "regioneProtectedAreas": {
                    "publisher": "Regione Toscana",
                    "page": "https://www.regione.toscana.it/-/il-sistema-delle-aree-naturali-protette",
                    "reference": "Archivi geografici ufficiali aree protette, Natura 2000 e Ramsar",
                    "layers": protected_provenance,
                },
                "regioneHydrography": {
                    "publisher": "Regione Toscana",
                    "page": "https://www.regione.toscana.it/-/reticolo-idrografico-e-di-gestione",
                    "reference": "Reticolo idrografico e di gestione · DCRT 24/2025",
                    "sourceLayer": reticulum_layer,
                    "acquisition": files["reticulum"],
                },
                "regioneIterNet": {
                    "publisher": "Regione Toscana",
                    "page": "https://www502.regione.toscana.it/geoscopio/download/grafo_stradale/",
                    "reference": "Iter.Net 4.48 · giugno 2022",
                    "sourceLayer": road_layer,
                    "acquisition": files["roads"],
                },
            },
            "measurementBoundary": {
                "publisher": "Istat", "reference": "Confini amministrativi al 1 gennaio 2026 · generalizzati",
                "sourceLayer": boundary_layer, "acquisition": files["boundaries"],
                "crs": "EPSG:3003", "istatCodes": TOWN_CODES,
            },
            "landUse": land_use,
            "landUseChange": land_change,
            "protectedNaturalAreas": protected,
            "reticulum": reticulum,
            "roadNetworkProfile": roads,
        }

    if set(payload["roadNetworkProfile"]["municipalities"]) != set(TOWNS):
        raise RuntimeError("Strade: copertura non 7/7")
    if payload["roadNetworkProfile"]["versilia"]["graphKm"] <= 0:
        raise RuntimeError("Strade: aggregato nullo")
    if set(payload["protectedNaturalAreas"]["municipalities"]) != set(TOWNS):
        raise RuntimeError("Aree protette: copertura non 7/7")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print("Snapshot v1.37 generato:", OUT)
    print("Reticolo gestito Versilia km:", payload["reticulum"]["versilia"]["managedNetworkKm"])
    print("Reticolo totale Versilia km:", payload["reticulum"]["versilia"]["fullNetworkKm"])
    print("Grafo viario Versilia km:", payload["roadNetworkProfile"]["versilia"]["graphKm"])
    print("Aree tutelate unione Versilia ha:", payload["protectedNaturalAreas"]["versilia"]["totalUnionHa"])
    for town in TOWNS:
        road = payload["roadNetworkProfile"]["municipalities"][town]
        hydro = payload["reticulum"]["municipalities"][town]
        prot = payload["protectedNaturalAreas"]["municipalities"][town]
        print(f"{town}: roads={road['graphKm']:.3f} km; hydro={hydro['fullNetworkKm']:.3f} km; protected={prot['totalUnionHa']:.2f} ha")


if __name__ == "__main__":
    main()
