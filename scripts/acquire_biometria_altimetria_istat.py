#!/usr/bin/env python3
"""Audita i due XLSX canonici Istat 2021 usati dalla Biometria comunale.

Scarica Altimetria e Fasce altimetriche, quindi isola intestazioni e righe dei sette
Comuni dell'Osservatorio. Non modifica i dati pubblicati.
"""
from __future__ import annotations

import argparse
import json
from io import BytesIO
from pathlib import Path

import requests
from openpyxl import load_workbook

SOURCES = {
    "altimetria": "https://www.istat.it/wp-content/uploads/2026/02/Altimetria_Comuni-al-31_12_2021.xlsx",
    "fasceAltimetriche": "https://www.istat.it/wp-content/uploads/2026/02/Fasce-altimetriche-Comuni-al-31_12_2021.xlsx",
}
TOWNS = ["Camaiore", "Forte dei Marmi", "Massarosa", "Pietrasanta", "Seravezza", "Stazzema", "Viareggio"]


def norm(value: object) -> str:
    return " ".join(str(value or "").strip().split()).casefold()


def audit_workbook(url: str) -> dict:
    response = requests.get(url, timeout=60)
    response.raise_for_status()
    workbook = load_workbook(BytesIO(response.content), data_only=True, read_only=True)
    wanted = {norm(name): name for name in TOWNS}
    matches: list[dict] = []
    sheet_heads: dict[str, list[list[object]]] = {}
    for ws in workbook.worksheets:
        head: list[list[object]] = []
        for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = list(row)
            if row_index <= 12:
                head.append(values)
            row_norm = {norm(value) for value in values if value is not None}
            hit = [canonical for key, canonical in wanted.items() if key in row_norm]
            if hit:
                matches.append({"sheet": ws.title, "row": row_index, "towns": hit, "values": values})
        sheet_heads[ws.title] = head
    missing = sorted(set(TOWNS) - {town for match in matches for town in match["towns"]})
    return {
        "source": url,
        "bytes": len(response.content),
        "sheets": workbook.sheetnames,
        "sheetHeads": sheet_heads,
        "matches": matches,
        "missing": missing,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", default="reports/biometria-comune-browser/istat-altimetria-raw-rows.json")
    args = parser.parse_args()

    payload = {key: audit_workbook(url) for key, url in SOURCES.items()}
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))

    missing = {key: value["missing"] for key, value in payload.items() if value["missing"]}
    if missing:
        raise SystemExit(f"Comuni non trovati nei file Istat: {missing}")


if __name__ == "__main__":
    main()
