"""Fixed ACI cells, raw workbook replay, public precision and adversarial guards."""
import base64
import io
import math

from semantic_query_engine import QueryEngine, ROOT
from semantic_question_suite import pointer
from test_semantic_query_engine import rejected
import semantic_query_vehicle_adapters as v

# Independently transcribed seven municipality records; total cars, Euro 0–3,
# POSAS residents 2026, published motorization, published Euro share.
FIXED = {'046005':(22650,4384,31763,713.1,19.4), '046013':(4839,939,6550,738.8,19.4),
         '046018':(16170,3062,21782,742.4,18.9), '046024':(16516,3308,22678,728.3,20.),
         '046028':(8607,1585,12284,700.7,18.4), '046030':(2060,509,2783,740.2,24.7),
         '046033':(39446,7158,60680,650.1,18.1)}
BENCHMARKS = {v.MOTOR:{'tuscany':744.1215646386034,'italy':701.366347742935},
              v.EURO:{'tuscany':17.166108561955497,'italy':23.822150647563276}}


def query(key, dim='total', op='compare', **selection):
    return dict(operation=op,selectors=[dict(metric=key,dimension=dim,**selection)])


def replay_workbook(snapshot):
    from openpyxl import load_workbook
    import hashlib
    body=base64.b64decode(snapshot['source']['workbookBase64'],validate=True)
    assert hashlib.sha256(body).hexdigest() == '700e7fbc0a1f3d68502aec979a11fc6d25f4c07ecc563fc356ee9d610b77dabc'
    book=load_workbook(io.BytesIO(body),read_only=True,data_only=True)
    sheet=book[snapshot['source']['sheet']];wanted={r[0]:r[1:] for r in snapshot['records']};checked=0;region=province=None
    for i,row in enumerate(sheet.iter_rows(values_only=True),1):
        if row[0] is not None: region=row[0]
        if row[1] is not None: province=row[1]
        if i in wanted:
            parsed=[region,province,*row[2:19]]
            assert parsed == wanted[i], (i,parsed,wanted[i])
            checked+=1
    book.close();assert checked == 7997


def regressions(path):
    e=QueryEngine(path,layer='effective');catalog=e.catalog
    snap,_=e.file(v.NATIVE);replay_workbook(snap)
    # Existing source audit independently reconciles native partition/geographic
    # sums and both public municipal rates; no network reacquisition.
    from acquire_a3_benchmark_aci_istat_annual import validate_snapshot
    for key in v.KEYS: validate_snapshot(catalog['metrics'][key],snap,catalog['metrics']['population'])
    for key in v.KEYS:
        j=3 if key == v.MOTOR else 4; scale=1000 if key == v.MOTOR else 100
        native={c:(x[0]/x[2]*1000 if key == v.MOTOR else x[1]/x[0]*100) for c,x in FIXED.items()}
        for dim in v.dimensions(key):
            r=e.query(query(key,dim));assert r['status']=='computed',r
            assert r['coverage']['requested']==7 and r['coverage']['usable']==7
            for o in r['observations']:
                code=o['geography'];x=FIXED[code];want=x[j] if dim=='total' else native[code]
                assert math.isclose(o['value'],want,rel_tol=0,abs_tol=1e-12)
                assert o['period']=='2024' and o['unit']==v.UNITS[key]
                assert pointer(catalog,o['provenance'][0]['valuePointer'])==x[j]
                record=pointer(snap,o['provenance'][1]['recordPointer']);assert record[-1]==x[0] and sum(record[4:8])==x[1]
                if dim=='nativeRatio':
                    assert o['numerator']==(x[0] if key==v.MOTOR else x[1]) and o['denominator']==(x[2] if key==v.MOTOR else x[0])
                    assert o['scale']==scale and o['denominatorPeriod']==('2026-01-01' if key==v.MOTOR else '2024')
                else:assert 'numerator' not in o
                if key==v.MOTOR:
                    pop,_=e.file(v.POPULATION);assert pointer(pop,o['provenance'][2]['recordPointer'])==x[2]
            ranked=e.query(query(key,dim,'rank'));assert ranked['status']=='computed'
            values={c:x[j] for c,x in FIXED.items()} if dim=='total' else native
            assert [x['geography'] for x in ranked['result']['ranking']]==sorted(values,key=lambda c:(-values[c],c))
            for scope,b in BENCHMARKS[key].items():
                for code in FIXED:
                    gap=e.query(dict(query(key,dim,'benchmark_gap',towns=[code]),benchmark=scope));assert gap['status']=='computed',gap
                    assert math.isclose(gap['result']['value'],values[code]-b,rel_tol=0,abs_tol=1e-10)
            rejected(e,dict(query(key,dim,'benchmark_gap',towns=['046018']),benchmark='versilia'),'vehicles_benchmark_scope_not_reviewed')
            for op in ('series','absolute_change','relative_change','percentage_points','trend'):
                rejected(e,query(key,dim,op,towns=['046018']),'vehicles_single_snapshot_no_history')
            rejected(e,query(key,dim,'anomaly'),'vehicles_peer_anomaly_not_reviewed')
            rejected(e,query(key,dim,periods=['2025']),'vehicles_period_not_frozen')
            for order in ([key,'population'],['population',key]):
                rejected(e,dict(operation='correlation',selectors=[dict(metric=k,dimension=dim if k==key else 'total') for k in order],axis='municipalities',method='pearson',purpose='descriptive'),'vehicles_pair_not_jointly_reviewed')
        for order in ([key,'roadSafety'],['roadSafety',key]):
            pair=e.query(dict(operation='correlation',selectors=[dict(metric=k) for k in order],axis='municipalities',method='pearson',purpose='descriptive'))
            assert pair['status']=='not_computable' and any(x in pair['reasons'] for x in ('vehicles_pair_not_jointly_reviewed','road_pair_not_jointly_reviewed'))
        rejected(e,query(key,'total','weighted_ratio'),'vehicles_native_ratio_dimension_required')
        pooled=e.query(query(key,'nativeRatio','weighted_ratio'));assert pooled['status']=='computed',pooled
        num=sum(x[0 if key==v.MOTOR else 1] for x in FIXED.values());den=sum(x[2 if key==v.MOTOR else 0] for x in FIXED.values())
        assert pooled['result']['numerator']==num and pooled['result']['denominator']==den
        assert math.isclose(pooled['result']['value'],num/den*scale,abs_tol=1e-12)
        aggregate = num/den*scale if key==v.MOTOR else sum(x[0]*x[4] for x in FIXED.values())/sum(x[0] for x in FIXED.values())
        assert math.isclose(catalog['metrics'][key]['aggregate']['value'],aggregate,abs_tol=1e-12)
        if key==v.EURO:assert not math.isclose(aggregate,pooled['result']['value'],abs_tol=1e-6)
        rejected(e,query(key,'sex:women'),'dimension_adapter_not_implemented')
        changes=[lambda m:m['meta'].update(year='2026'),lambda m:m['meta'].update(theme='economia'),lambda m:m['meta'].update(unit='number'),lambda m:m['meta'].update(polarity='positive'),lambda m:m.update(sourceUrl='https://example.invalid'),lambda m:m['method'].update(formula='Other'),lambda m:m['method'].update(coverage='6/7'),lambda m:m['rows'][0].update(value=999),lambda m:m['rows'][0].update(value=False),lambda m:m['rows'][0].update(benchmarkValue=999),lambda m:m['rows'][0].update(town='Unknown'),lambda m:m['rows'][0].update(code='046005'),lambda m:m['rows'][0].update(slug='unknown'),lambda m:m['rows'][0].update(series={'years':[2024],'values':[1]}),lambda m:m['rows'][0].update(ratioComponents={'numerator':1}),lambda m:m['rows'][0].update(notApplicable=True),lambda m:m['rows'][0].update(dataUnavailable=True),lambda m:m['aggregate'].update(value=999),lambda m:m['aggregate'].update(label='Mean'),lambda m:m['meta']['benchmark'].update(tuscany=999),lambda m:m['meta']['benchmark'].update(sourceSnapshot=v.DEMO)]
        for change in changes:
            g=QueryEngine(path,layer='effective');change(g._catalog['metrics'][key]);assert g.query(query(key))['status']=='not_computable'
        for mutate in [lambda m:m['meta'].update(year='2024'),lambda m:m['rows'][0].update(value=0),lambda m:m['rows'][0].update(code='999999')]:
            g=QueryEngine(path,layer='effective');mutate(g._catalog['metrics']['population']);assert g.query(query(key))['status']=='not_computable'
    for path in v.HASHES:
        for field in ('sha256','path'):
            g=QueryEngine(e.catalog_path,layer='effective');_,ref=g.file(path);ref[field]='changed';assert g.query(query(v.MOTOR))['status']=='not_computable'
        g=QueryEngine(e.catalog_path,layer='effective');data,_=g.file(path);data['schemaVersion']=999;assert g.query(query(v.EURO))['status']=='not_computable'
    print('Vehicles PASS: 14 published and 14 distinct native ratios, 21 fixed components, 56 gaps, 2 pooled ratios, 7997 workbook localization rows replayed, identity/precision/date/adversarial guards')


if __name__=='__main__': regressions(ROOT/'dist/data/site-data.json')
