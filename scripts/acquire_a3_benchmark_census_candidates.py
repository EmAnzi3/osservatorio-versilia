#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
GOVERNED = ROOT / "data" / "source-snapshots" / "istat-sections-history-v1.8.0.json"

URL = "https://esploradati.istat.it/databrowser/DWL/PERMPOP/SUBCOM/Dati_regionali_2023.zip"
YEAR = "2023"

BASE_FIELDS = ["P1", "P102", "P103", "PF1", "PF3", "A3", "A8"]
MALE_FIELDS = [f"P{i}" for i in range(33, 43)]
FEMALE_FIELDS = [f"P{i}" for i in range(70, 80)]
REQUIRED_FIELDS = BASE_FIELDS + MALE_FIELDS + FEMALE_FIELDS

METRICS = {
    "femaleEmploymentRate": ("P103", "female1564", 100.0, "percent"),
    "maleEmploymentRate": ("P102", "male1564", 100.0, "percent"),
    "housingStockPer1000": ("A8", "P1", 1000.0, "per1000"),
    "nonOccupiedHomesPer1000": ("A3", "P1", 1000.0, "per1000"),
    "vacantHomes": ("A3", "A8", 100.0, "percent"),
    "singleHouseholds": ("PF3", "PF1", 100.0, "percent"),
}

MUNICIPALITY_HEADER_KEYS = {
    "PROCOM", "PROCOMT", "CODCOM", "CODICECOMUNE", "CODCOMUNE",
    "PROCOMUNE", "CODICEISTATCOMUNE",
}


def norm_header(value: Any) -> str:
    return re.sub(r"[^A-Z0-9]+", "", str(value or "").upper())


def code6(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return f"{value:06d}"
    if isinstance(value, float) and value.is_integer():
        return f"{int(value):06d}"
    s = str(value).strip()
    if s.endswith(".0"):
        s = s[:-2]
    digits = re.sub(r"\D", "", s)
    return digits.zfill(6) if 1 <= len(digits) <= 6 else ""


def number(value: Any, label: str) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        raise RuntimeError(f"{label}: booleano inatteso")
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        s = str(value).strip().replace(" ", "").replace(",", ".")
        try:
            result = float(s)
        except ValueError as exc:
            raise RuntimeError(f"{label}: valore non numerico {value!r}") from exc
    if not math.isfinite(result) or result < 0:
        raise RuntimeError(f"{label}: valore invalido {value!r}")
    return result


def find_header(ws) -> tuple[int, dict[str, int], list[Any]]:
    required = set(REQUIRED_FIELDS)
    for row_no, row in enumerate(ws.iter_rows(min_row=1, max_row=30, values_only=True), 1):
        normalized = [norm_header(v) for v in row]
        positions = {name: idx for idx, name in enumerate(normalized) if name}
        if required.issubset(positions):
            return row_no, positions, list(row)
    raise RuntimeError("Header censuario con variabili richieste non trovato nelle prime 30 righe")


def municipality_column(positions: dict[str, int]) -> int:
    candidates = [(name, idx) for name, idx in positions.items() if name in MUNICIPALITY_HEADER_KEYS]
    if len(candidates) == 1:
        return candidates[0][1]
    if candidates:
        # PRO_COM è il contratto storico più comune; preferiscilo se presente.
        for name, idx in candidates:
            if name == "PROCOM":
                return idx
    raise RuntimeError(f"Colonna codice Comune non individuata; header candidati={candidates}")


def empty_totals() -> dict[str, float]:
    return {field: 0.0 for field in REQUIRED_FIELDS}


def add_row(target: dict[str, float], row: tuple[Any, ...], positions: dict[str, int], label: str) -> None:
    for field in REQUIRED_FIELDS:
        idx = positions[field]
        value = row[idx] if idx < len(row) else None
        target[field] += number(value, f"{label}/{field}")


def finalize(raw: dict[str, float]) -> dict[str, float]:
    out = {k: float(v) for k, v in raw.items()}
    out["male1564"] = sum(out[f"P{i}"] for i in range(33, 43))
    out["female1564"] = sum(out[f"P{i}"] for i in range(70, 80))
    return out


def parse_workbook(blob: bytes, member: str, town_codes: set[str] | None = None):
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        header_row, positions, headers = find_header(ws)
        mun_idx = municipality_column(positions)

        total = empty_totals()
        towns: dict[str, dict[str, float]] = defaultdict(empty_totals)
        rows = 0
        matched_rows = 0

        for row_no, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            if not row or all(v in (None, "") for v in row):
                continue
            add_row(total, row, positions, f"{member}:r{row_no}")
            rows += 1
            if town_codes:
                code = code6(row[mun_idx] if mun_idx < len(row) else None)
                if code in town_codes:
                    add_row(towns[code], row, positions, f"{member}:{code}:r{row_no}")
                    matched_rows += 1

        if rows <= 0:
            raise RuntimeError(f"{member}: nessuna riga dati")
        return finalize(total), {k: finalize(v) for k, v in towns.items()}, {
            "sheet": ws.title,
            "headerRow": header_row,
            "rows": rows,
            "matchedTownRows": matched_rows,
            "municipalityHeader": headers[mun_idx] if mun_idx < len(headers) else None,
            "headersRequired": {field: headers[positions[field]] for field in REQUIRED_FIELDS},
        }
    finally:
        wb.close()


def ratio(raw: dict[str, float], numerator: str, denominator: str, scale: float) -> float:
    den = raw[denominator]
    if den <= 0:
        raise RuntimeError(f"Denominatore nullo per {numerator}/{denominator}")
    return raw[numerator] / den * scale


def compare_governed(towns: dict[str, dict[str, float]], governed: dict[str, Any]) -> dict[str, Any]:
    expected_rows = {
        str(row["code"]): row
        for row in ((governed.get("raw") or {}).get(YEAR) or [])
        if isinstance(row, dict) and row.get("code")
    }
    errors = []
    checks = ["P1", "P102", "P103", "PF1", "PF3", "A3", "A8", "male1564", "female1564"]
    for code, expected in expected_rows.items():
        got = towns.get(code)
        if got is None:
            errors.append(f"{code}: Comune non ricostruito")
            continue
        for field in checks:
            actual = got[field]
            wanted = float(expected[field])
            if not math.isclose(actual, wanted, rel_tol=0.0, abs_tol=1e-9):
                errors.append(f"{code}/{field}: {actual} != {wanted}")
    extras = sorted(set(towns) - set(expected_rows))
    if extras:
        errors.append(f"Codici Versilia inattesi: {extras}")
    return {
        "status": "PASS" if not errors and len(towns) == len(expected_rows) == 7 else "FAIL",
        "townsExpected": len(expected_rows),
        "townsReconstructed": len(towns),
        "errors": errors[:100],
    }


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--output", required=True)
    args = ap.parse_args()

    governed = json.loads(GOVERNED.read_text(encoding="utf-8"))
    town_codes = {
        str(item["code"])
        for item in ((governed.get("scope") or {}).get("towns") or [])
        if isinstance(item, dict) and item.get("code")
    }
    if len(town_codes) != 7:
        raise RuntimeError("Perimetro governato Versilia non 7/7")

    session = requests.Session()
    session.headers["User-Agent"] = "OsservatorioVersilia-A3-census-benchmark/2.0"

    with tempfile.NamedTemporaryFile(suffix=".zip") as tmp:
        sha = hashlib.sha256()
        total_bytes = 0
        with session.get(URL, stream=True, timeout=240) as response:
            response.raise_for_status()
            for chunk in response.iter_content(chunk_size=1024 * 1024):
                if not chunk:
                    continue
                tmp.write(chunk)
                sha.update(chunk)
                total_bytes += len(chunk)
        tmp.flush()

        italy_acc = empty_totals()
        tuscany = None
        tuscany_towns = {}
        diagnostics = {}
        workbook_hashes = {}

        with zipfile.ZipFile(tmp.name) as archive:
            members = [
                name for name in archive.namelist()
                if re.search(r"R\d{2}_.+_2023_sezioni\.xlsx$", Path(name).name, re.IGNORECASE)
                and not Path(name).name.startswith("~$")
            ]
            if not (20 <= len(members) <= 22):
                raise RuntimeError(f"Numero file regionali inatteso: {len(members)}; sample={members[:10]}")

            region_codes = []
            for member in sorted(members):
                base = Path(member).name
                match = re.match(r"R(\d{2})_", base, re.IGNORECASE)
                if not match:
                    raise RuntimeError(f"Codice regione non ricavabile da {base}")
                region_code = match.group(1)
                region_codes.append(region_code)

                blob = archive.read(member)
                workbook_hashes[base] = hashlib.sha256(blob).hexdigest()
                is_tuscany = region_code == "09"
                regional, towns, diag = parse_workbook(blob, base, town_codes if is_tuscany else None)
                diagnostics[base] = diag
                for field in REQUIRED_FIELDS:
                    italy_acc[field] += regional[field]
                if is_tuscany:
                    tuscany = regional
                    tuscany_towns = towns

            if len(set(region_codes)) != len(region_codes):
                raise RuntimeError(f"Codici regione duplicati nel pacchetto: {region_codes}")

    if tuscany is None:
        raise RuntimeError("File Toscana R09 non trovato")

    italy = finalize(italy_acc)
    reconcile = compare_governed(tuscany_towns, governed)

    sanity = {
        "regionalFiles": 20 <= len(diagnostics) <= 22,
        "tuscanyPopulation": 3_000_000 <= tuscany["P1"] <= 4_500_000,
        "italyPopulation": 55_000_000 <= italy["P1"] <= 65_000_000,
        "tuscanyVsItaly": 0 < tuscany["P1"] < italy["P1"],
        "versiliaReconciliation": reconcile["status"] == "PASS",
    }

    raw = {"tuscany": tuscany, "italy": italy}
    benchmarks = {}
    for metric_id, (num, den, scale, unit) in METRICS.items():
        benchmarks[metric_id] = {
            "year": 2023,
            "unit": unit,
            "formula": f"{num} / {den} × {scale:g}",
            "tuscany": ratio(tuscany, num, den, scale),
            "italy": ratio(italy, num, den, scale),
            "raw": {
                "tuscany": {"numerator": tuscany[num], "denominator": tuscany[den]},
                "italy": {"numerator": italy[num], "denominator": italy[den]},
            },
        }

    status = "ACQUIRED_CANDIDATE" if all(sanity.values()) else "CANDIDATE_REJECTED"
    payload = {
        "schemaVersion": 2,
        "publisher": "Istat — Censimento permanente della popolazione e delle abitazioni",
        "profileId": "istat-census-annual",
        "referenceYear": 2023,
        "source": {
            "url": URL,
            "bytes": total_bytes,
            "sha256": sha.hexdigest(),
            "regionalWorkbookCount": len(diagnostics),
            "regionalWorkbookHashes": workbook_hashes,
        },
        "method": "Aggregazione delle variabili additive nei file regionali ufficiali per sezione 2023. Toscana = file R09; Italia = somma dei file regionali. Nessuna derivazione dal perimetro subcomunale Comuni_2023.",
        "raw": raw,
        "benchmarks": benchmarks,
        "tuscanyVersiliaReconciliation": reconcile,
        "diagnostics": diagnostics,
        "sanityGate": sanity,
        "status": status,
    }

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "status": status,
        "metrics": sorted(benchmarks),
        "tuscanyPopulation": tuscany["P1"],
        "italyPopulation": italy["P1"],
        "reconciliation": reconcile["status"],
        "output": str(out),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
