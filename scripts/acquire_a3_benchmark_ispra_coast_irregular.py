#!/usr/bin/env python3
from __future__ import annotations
import argparse,io,json,math
from pathlib import Path
from typing import Any
import requests
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
SITE=ROOT/"data/site-data.json"
SNAP=ROOT/"data/source-snapshots/costa-mare-v123.json"
URL="https://indicatoriambientali.isprambiente.it/sites/default/files/indicatori_ambientali/costa-protetta/01_TB1_CProtetta2020%20%28rev%29.xlsx"
LANDING="https://indicatoriambientali.isprambiente.it/it/coste/costa-protetta"
COASTAL={"046005","046013","046024","046033"}
NA={"046018","046028","046030"}

def finite(v:Any)->float:
    if isinstance(v,bool): raise ValueError(v)
    x=float(v)
    if not math.isfinite(x): raise ValueError(v)
    return x

def share(n:float,d:float)->float:
    if d<=0: raise RuntimeError("denominatore costa non positivo")
    return n/d*100.0

def main()->None:
    ap=argparse.ArgumentParser(); ap.add_argument("--output",required=True); a=ap.parse_args()
    site=json.loads(SITE.read_text(encoding="utf-8"))
    snap=json.loads(SNAP.read_text(encoding="utf-8"))
    metric=(site.get("metrics") or {}).get("rigidDefenceProtectedCoast") or {}
    rows=metric.get("rows") or []
    if len(rows)!=7: raise RuntimeError(f"costa protetta: pubblico {len(rows)}/7")
    governed=(snap.get("rigidDefenceProtectedCoast2020") or {}).get("towns") or {}
    if set(governed)!=COASTAL: raise RuntimeError("costa protetta: snapshot costiero inatteso")
    errors=[]
    for row in rows:
        code=str(row.get("code") or ""); value=row.get("value")
        if code in NA:
            if value is not None: errors.append(f"{code}: atteso n.a., trovato {value}")
            continue
        if code not in COASTAL:
            errors.append(f"{code}: codice fuori perimetro"); continue
        raw=governed[code]
        expected=share(finite(raw["protectedKm"]),finite(raw["coastKm"]))
        if value is None or not math.isclose(float(value),expected,rel_tol=0.0,abs_tol=.011):
            errors.append(f"{code}: {value} != {expected}")

    r=requests.get(URL,timeout=180,headers={"User-Agent":"OsservatorioVersilia-A3-coast-benchmark/1.0"})
    r.raise_for_status()
    wb=load_workbook(io.BytesIO(r.content),read_only=True,data_only=True)
    ws=wb["dati"] if "dati" in wb.sheetnames else wb[wb.sheetnames[0]]
    found={}
    for row in ws.iter_rows(values_only=True):
        name=str(row[0] or "").strip().casefold()
        if name in {"toscana","italia"}:
            coast=finite(row[7]); protected=finite(row[8]); pct=finite(row[9])
            calc=share(protected,coast)
            if not math.isclose(calc,pct,rel_tol=0.0,abs_tol=1e-6):
                raise RuntimeError(f"ISPRA costa {name}: {calc} != {pct}")
            found[name]={"coastKm":coast,"protectedKm":protected,"percent":pct}
    if set(found)!={"toscana","italia"}:
        raise RuntimeError(f"ISPRA costa: righe Toscana/Italia non univoche {found}")

    gate="PASS" if not errors else "FAIL"
    payload={
      "schemaVersion":2,"publisher":"ISPRA — Costa protetta","profileId":"ispra-coast-irregular",
      "scope":{"coastalTownCodes":sorted(COASTAL),"notApplicableTownCodes":sorted(NA),"reconciled":"4 coastal + 3 n.a."},
      "status":"ACQUIRED_CANDIDATE" if gate=="PASS" else "CANDIDATE_REJECTED",
      "sourceUrl":LANDING,"dataUrl":URL,
      "benchmarks":{"rigidDefenceProtectedCoast":{
        "year":"2020","unit":"percent",
        "formula":"km di costa protetta da opere rigide / km di costa dell'universo ISPRA × 100",
        "tuscany":found["toscana"]["percent"],"italy":found["italia"]["percent"],
      }} if gate=="PASS" else {},
      "components":found,
      "qualityGate":{"status":gate,"publicReconciliation":"1 metric × 4 coastal towns + 3 n.a. PASS" if gate=="PASS" else "FAIL","regionalNationalControl":"official 2020 km + percent rows PASS","errors":errors},
      "blocked":{"shorelineDynamics":"metrica composita erosione/stabile/avanzamento; benchmark scalare non sufficiente"}
    }
    out=Path(a.output); out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"status":payload["status"],"candidateMetrics":sorted(payload["benchmarks"]),"blocked":payload["blocked"]},ensure_ascii=False))
if __name__=="__main__": main()
