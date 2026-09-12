#!/usr/bin/env python3
"""Inventario diagnostico delle fonti ufficiali necessarie alla v1.37.

Non modifica data/site-data.json. Serve esclusivamente nella Action di acquisizione
per identificare struttura e URL degli input ufficiali prima di generare lo snapshot.
"""
from __future__ import annotations

import io
import json
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

from openpyxl import load_workbook

OUT = Path("tmp/territorio-v137-source-inventory.json")
UA = {"User-Agent": "OsservatorioVersilia-v137-source-audit/1.0"}
CKAN = "https://dati.toscana.it/api/3/action/package_show?id={}"
PACKAGES = {
    # I metadati CKAN sono storici, ma l'URL Geoscopio punta al pacchetto corrente.
    "roadsLegacyCatalog": "grafo-civici",
    "naturaLegacyCatalog": "sir",
    "ramsarLegacyCatalog": "ramsar",
}
DIRECT = {
    "ispraSoil": "https://www.isprambiente.gov.it/it/attivita/suolo-e-territorio/suolo/il-consumo-di-suolo/consumo_di_suolo_estratto_dati_2025_anni_2006_2024.xlsx",
    "istatBoundaries2026": "https://www.istat.it/storage/cartografia/confini_amministrativi/non_generalizzati/2026/Limiti01012026.zip",
    "reticulumPage": "https://www.regione.toscana.it/-/reticolo-idrografico-e-di-gestione",
    "reticulumZip": "https://www.regione.toscana.it/documents/d/guest/infrastruttura_rev25-zip",
    "roadsMetadata": "https://www502.regione.toscana.it/geonetwork/srv/search?keyword=sistema%20di%20trasporto",
    "roadsDownloadPage": "https://www502.regione.toscana.it/geoscopio/download/grafo_stradale/",
    "roadsZip": "https://www502.regione.toscana.it/geoscopio/download/grafo_stradale/iternet.zip",
    "protectedWmsCard": "https://www502.regione.toscana.it/geoscopio/servizi/wms/AREE_PROTETTE.htm",
    "natura2000Page": "https://www.regione.toscana.it/-/rete-natura-2000-in-toscana-2",
    "ramsarPage": "https://www.regione.toscana.it/-/aree-ramsar",
}
WFS_BASE = "https://www502.regione.toscana.it/wmsraster/com.rt.wms.RTmap/wms"
WFS_LAYERS = {
    "parksNational": "rt_arprot.idparnaz.rt.poly",
    "reservesNational": "rt_arprot.idrisnatstat.rt.poly",
    "parksRegional": "rt_arprot.idparreg.rt.poly",
    "parksProvincial": "rt_arprot.idparprov.rt.poly",
    "reservesRegional": "rt_arprot.idrisnatreg.rt.poly",
    "anpil": "rt_arprot.idanpil.rt.poly",
    "zsc": "rt_arprot.idnat2000_sic.rt.poly",
    "zps": "rt_arprot.idnat2000_zps.rt.poly",
    "zscZps": "rt_arprot.idnat2000_sic_zps.rt.poly",
    "ramsar": "rt_arprot.idramsar.rt.poly",
}


def request_bytes(url: str, timeout: int = 120) -> bytes:
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=timeout) as response:
        return response.read()


def request_json(url: str) -> dict:
    return json.loads(request_bytes(url, 60))


def probe(url: str) -> dict:
    req = urllib.request.Request(url, method="HEAD", headers=UA)
    try:
        with urllib.request.urlopen(req, timeout=30) as response:
            return {
                "status": response.status,
                "finalUrl": response.geturl(),
                "contentType": response.headers.get("Content-Type"),
                "contentLength": response.headers.get("Content-Length"),
                "etag": response.headers.get("ETag"),
                "lastModified": response.headers.get("Last-Modified"),
            }
    except urllib.error.HTTPError as exc:
        return {"status": exc.code, "error": str(exc)}
    except Exception as exc:  # noqa: BLE001
        return {"status": None, "error": repr(exc)}


def normalized_resource(resource: dict) -> dict:
    return {
        "id": resource.get("id"), "name": resource.get("name"), "format": resource.get("format"),
        "url": resource.get("url"), "size": resource.get("size"), "last_modified": resource.get("last_modified"),
        "created": resource.get("created"), "description": resource.get("description"),
    }


def zip_inventory(url: str) -> dict:
    raw = request_bytes(url, 240)
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        names = archive.namelist()
    return {"bytes": len(raw), "members": names}


def xlsx_inventory(url: str) -> dict:
    raw = request_bytes(url, 120)
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    sheets = {}
    for worksheet in workbook.worksheets:
        preview = []
        for row in worksheet.iter_rows(min_row=1, max_row=12, values_only=True):
            preview.append([value for value in row[:30]])
        sheets[worksheet.title] = {"maxRow": worksheet.max_row, "maxColumn": worksheet.max_column, "preview": preview}
    return {"bytes": len(raw), "sheets": sheets}


def wfs_url(layer: str) -> str:
    params = {
        "map": "wmsarprot", "service": "WFS", "version": "2.0.0", "request": "GetFeature",
        "OUTPUTFORMAT": "application/shapefile", "filename": f"v137_{layer.split('.')[-2]}", "typename": layer,
    }
    return WFS_BASE + "?" + urllib.parse.urlencode(params)


def main() -> None:
    output = {"packages": {}, "direct": {}, "archives": {}, "ispraWorkbook": {}, "protectedWfs": {}}
    failures = []

    for key, package_id in PACKAGES.items():
        try:
            payload = request_json(CKAN.format(package_id))
            result = payload["result"] if payload.get("success") else None
            if not result:
                raise RuntimeError("CKAN success=false")
            output["packages"][key] = {
                "packageId": package_id, "title": result.get("title"), "metadataModified": result.get("metadata_modified"),
                "resources": [normalized_resource(item) for item in result.get("resources", [])],
            }
        except Exception as exc:  # noqa: BLE001
            output["packages"][key] = {"packageId": package_id, "error": repr(exc)}

    for key, url in DIRECT.items():
        output["direct"][key] = {"url": url, "probe": probe(url)}

    for key in ("istatBoundaries2026", "reticulumZip", "roadsZip"):
        try:
            output["archives"][key] = zip_inventory(DIRECT[key])
        except Exception as exc:  # noqa: BLE001
            output["archives"][key] = {"error": repr(exc)}
            failures.append(f"{key}: {exc!r}")

    try:
        output["ispraWorkbook"] = xlsx_inventory(DIRECT["ispraSoil"])
    except Exception as exc:  # noqa: BLE001
        output["ispraWorkbook"] = {"error": repr(exc)}
        failures.append(f"ISPRA: {exc!r}")

    for key, layer in WFS_LAYERS.items():
        url = wfs_url(layer)
        try:
            raw = request_bytes(url, 180)
            content = {"url": url, "bytes": len(raw), "zip": False}
            if raw[:2] == b"PK":
                with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                    content.update({"zip": True, "members": archive.namelist()})
            else:
                content["prefix"] = raw[:500].decode("utf-8", errors="replace")
            output["protectedWfs"][key] = content
            if not content["zip"]:
                failures.append(f"WFS {key}: risposta non shapefile ZIP")
        except Exception as exc:  # noqa: BLE001
            output["protectedWfs"][key] = {"url": url, "error": repr(exc)}
            failures.append(f"WFS {key}: {exc!r}")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(output, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    print(json.dumps(output, ensure_ascii=False, indent=2, default=str))
    if failures:
        raise SystemExit("Inventario incompleto: " + "; ".join(failures))


if __name__ == "__main__":
    main()
