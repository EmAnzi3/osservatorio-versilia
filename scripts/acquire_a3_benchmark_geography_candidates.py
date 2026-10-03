#!/usr/bin/env python3
from __future__ import annotations
import argparse, io, json, math, re
from pathlib import Path
from typing import Any
import requests
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
SNAP=ROOT/'data/source-snapshots/biometria-comune-v135.json'
DEMO=ROOT/'data/source-snapshots/a3-istat-demography-benchmark-2026.json'
ALT_URL='https://www.istat.it/wp-content/uploads/2026/02/Altimetria_Comuni-al-31_12_2021.xlsx'
BAND_URL='https://www.istat.it/wp-content/uploads/2026/02/Fasce-altimetriche-Comuni-al-31_12_2021.xlsx'
TUSCANY_PREFIXES={'045','046','047','048','049','050','051','052','053','100'}
TOWNS={'046005':'Camaiore','046013':'Forte dei Marmi','046018':'Massarosa','046024':'Pietrasanta','046028':'Seravezza','046030':'Stazzema','046033':'Viareggio'}

def norm(v:Any)->str: return re.sub(r'[^A-Z0-9<>]+','',str(v or '').upper())
def code6(v:Any)->str:
    s=re.sub(r'\D','',str(v or '').strip().removesuffix('.0'))
    return s.zfill(6) if 1<=len(s)<=6 else ''
def number(v:Any,label:str)->float:
    if isinstance(v,(int,float)) and not isinstance(v,bool): x=float(v)
    else: x=float(str(v or '').strip().replace('.','').replace(',','.'))
    if not math.isfinite(x): raise RuntimeError(f'{label}: non finito')
    return x
def workbook(url:str):
    r=requests.get(url,timeout=180,headers={'User-Agent':'OsservatorioVersilia-A3-geography/1.0'}); r.raise_for_status()
    return load_workbook(io.BytesIO(r.content),read_only=True,data_only=True),r.content
def header(ws, required:set[str]):
    for rn,row in enumerate(ws.iter_rows(min_row=1,max_row=20,values_only=True),1):
        vals=list(row); pos={norm(v):i for i,v in enumerate(vals) if norm(v)}
        if required.issubset(pos): return rn,pos,vals
    raise RuntimeError(f'{ws.title}: header non trovato {required}')
def band_id(raw:Any)->str|None:
    s=str(raw or '').lower().replace(' ','').replace('.','')
    if 'perc' not in s: return None
    if '<300' in s: return '0_299'
    if '300' in s and '599' in s: return '300_599'
    if '600' in s and '899' in s: return '600_899'
    if '900' in s and '1199' in s: return '900_1199'
    if '1200' in s and '1499' in s: return '1200_1499'
    if '1500' in s and '1999' in s: return '1500_1999'
    if '2000' in s and '2499' in s: return '2000_2499'
    if '>2500' in s or '2500+' in s: return '2500_plus'
    return None
def parse_alt(wb):
    out={}
    for ws in wb.worksheets:
        try: rn,pos,_=header(ws,{'PROCOM','AREAKMQ'})
        except RuntimeError: continue
        for row in ws.iter_rows(min_row=rn+1,values_only=True):
            code=code6(row[pos['PROCOM']] if pos['PROCOM']<len(row) else None)
            if not re.fullmatch(r'\d{6}',code): continue
            area=number(row[pos['AREAKMQ']],f'{code}/area')
            if area<=0: continue
            if code in out and not math.isclose(out[code],area,abs_tol=1e-9): raise RuntimeError(f'{code}: area duplicata incoerente')
            out[code]=area
    return out
def parse_bands(wb):
    out={}
    for ws in wb.worksheets:
        try: rn,pos,heads=header(ws,{'PROCOM','AREAHA'})
        except RuntimeError: continue
        bcols={band_id(h):i for i,h in enumerate(heads) if band_id(h)}
        if len(bcols)!=8: continue
        for row in ws.iter_rows(min_row=rn+1,values_only=True):
            code=code6(row[pos['PROCOM']] if pos['PROCOM']<len(row) else None)
            if not re.fullmatch(r'\d{6}',code): continue
            area=number(row[pos['AREAHA']],f'{code}/areaHa')
            if area<=0: continue
            vals={k:number(row[i],f'{code}/{k}') for k,i in bcols.items()}
            if not math.isclose(sum(vals.values()),100.0,abs_tol=.2): raise RuntimeError(f'{code}: fasce sommano {sum(vals.values())}')
            out[code]={'areaHa':area,'bands':vals}
    return out
def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--output',required=True); a=ap.parse_args()
    snap=json.loads(SNAP.read_text(encoding='utf-8')); demo=json.loads(DEMO.read_text(encoding='utf-8'))
    alt_wb,alt_blob=workbook(ALT_URL); band_wb,band_blob=workbook(BAND_URL)
    try: areas=parse_alt(alt_wb); bands=parse_bands(band_wb)
    finally: alt_wb.close(); band_wb.close()
    common=set(areas)&set(bands)
    if not 7800<=len(common)<=8000: raise RuntimeError(f'perimetro nazionale inatteso {len(common)}')
    tus={c for c in common if c[:3] in TUSCANY_PREFIXES}
    if not 270<=len(tus)<=280: raise RuntimeError(f'perimetro Toscana inatteso {len(tus)}')
    errors=[]
    for code,name in TOWNS.items():
        src=snap['municipalities'][name]
        if code not in common: errors.append(f'{name}: assente'); continue
        if not math.isclose(areas[code],float(src['surfaceKm2']),abs_tol=1e-6): errors.append(f'{name}: area {areas[code]} != {src["surfaceKm2"]}')
        for k,v in src['altitudeBandsPct'].items():
            if not math.isclose(bands[code]['bands'][k],float(v),abs_tol=.05): errors.append(f'{name}/{k}: {bands[code]["bands"][k]} != {v}')
    def scope(codes:set[str],pop:float):
        area=sum(areas[c] for c in codes); area_ha=sum(bands[c]['areaHa'] for c in codes)
        by={k:sum(bands[c]['areaHa']*bands[c]['bands'][k]/100 for c in codes) for k in next(iter(bands.values()))['bands']}
        return {'municipalities':len(codes),'totalSurfaceKm2':area,'meanMunicipalSurfaceKm2':area/len(codes),'population':pop,'populationDensity':pop/area,'altitudeBandsPct':{k:v/area_ha*100 for k,v in by.items()},'from300Pct':sum(v for k,v in by.items() if k!='0_299')/area_ha*100}
    scopes={'tuscany':scope(tus,float(demo['benchmarks']['population']['tuscany'])),'italy':scope(common,float(demo['benchmarks']['population']['italy']))}
    benchmarks={
      'municipalSurface':{'year':'31 dicembre 2021','unit':'squareKm','mapping':'mean municipal surface within scope','tuscany':scopes['tuscany']['meanMunicipalSurfaceKm2'],'italy':scopes['italy']['meanMunicipalSurfaceKm2']},
      'populationDensity':{'year':'2026 / superficie 2021','unit':'peoplePerSquareKm','mapping':'scope population / scope surface','tuscany':scopes['tuscany']['populationDensity'],'italy':scopes['italy']['populationDensity']},
      'altitudeProfile':{'year':'31 dicembre 2021','unit':'percent','part':'Territorio da 300 m in su','mapping':'area-weighted share >=300 m','tuscany':scopes['tuscany']['from300Pct'],'italy':scopes['italy']['from300Pct']}
    }
    gate={'status':'PASS' if not errors else 'FAIL','municipalityCountItaly':len(common),'municipalityCountTuscany':len(tus),'publicSnapshotReconciliation':'7/7 PASS' if not errors else 'FAIL','errors':errors}
    payload={'schemaVersion':1,'publisher':'Istat','profileId':'istat-geografia-comunale-2021','referenceYear':2021,'status':'ACQUIRED_CANDIDATE' if not errors else 'CANDIDATE_REJECTED','sources':{'altimetry':ALT_URL,'bands':BAND_URL,'altimetryBytes':len(alt_blob),'bandsBytes':len(band_blob)},'method':'Aggregazione dei Comuni italiani canonici a 6 cifre. Superficie benchmark = superficie media comunale dello scope; densità = popolazione aggregata / superficie aggregata; altimetria = quota di area >=300 m ponderata per superficie.','benchmarks':benchmarks,'scopes':scopes,'qualityGate':gate}
    p=Path(a.output); p.parent.mkdir(parents=True,exist_ok=True); p.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'status':payload['status'],'metrics':list(benchmarks),'gate':gate},ensure_ascii=False))
if __name__=='__main__': main()
