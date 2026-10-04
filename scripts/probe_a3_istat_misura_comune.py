#!/usr/bin/env python3
from __future__ import annotations

import argparse
import io
import json
from pathlib import Path

import requests
from openpyxl import load_workbook

TARGETS = (
    "camaiore",
    "forte dei marmi",
    "massarosa",
    "pietrasanta",
    "seravezza",
    "stazzema",
    "viareggio",
    "toscana",
    "italia",
)


def clean_row(values):
    result = []
    for value in values:
        if value is None:
            result.append(None)
        elif isinstance(value, str):
            result.append(value.strip())
        else:
            result.append(value)
    while result and result[-1] in (None, ""):
        result.pop()
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--slug", required=True)
    parser.add_argument("--url", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    session = requests.Session()
    session.headers["User-Agent"] = "OsservatorioVersilia-A3-Istat-MisuraComune/1.0"
    response = session.get(args.url, timeout=240)
    response.raise_for_status()
    body = response.content

    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    sheets = []
    total_matches = 0
    for sheet in workbook.worksheets:
        header_samples = []
        matches = []
        nonempty_seen = 0
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = clean_row(list(row))
            if not any(value not in (None, "") for value in values):
                continue
            if nonempty_seen < 18:
                header_samples.append({"row": row_index, "values": values[:80]})
                nonempty_seen += 1
            text = " | ".join(str(value) for value in values if value not in (None, "")).lower()
            found = [target for target in TARGETS if target in text]
            if found:
                matches.append({
                    "row": row_index,
                    "targets": found,
                    "values": values[:120],
                })
                total_matches += 1
        sheets.append({
            "title": sheet.title,
            "maxRow": sheet.max_row,
            "maxColumn": sheet.max_column,
            "headerSamples": header_samples,
            "matches": matches,
        })

    report = {
        "schemaVersion": 1,
        "slug": args.slug,
        "url": args.url,
        "resolvedUrl": response.url,
        "bytes": len(body),
        "sheetCount": len(sheets),
        "matchCount": total_matches,
        "sheets": sheets,
    }
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "slug": args.slug,
        "bytes": len(body),
        "sheets": len(sheets),
        "matches": total_matches,
        "output": str(output),
    }, ensure_ascii=False))
    if total_matches == 0:
        raise SystemExit(f"{args.slug}: nessuna riga Versilia/Toscana/Italia trovata")


if __name__ == "__main__":
    main()
