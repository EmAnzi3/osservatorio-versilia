#!/usr/bin/env python3
from __future__ import annotations

import argparse
import io
import json
import math
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
SITE=ROOT/"data/site-data.json"
MUNICIPAL_URL="https://www.istat.it/storage/misura-comune/10b-Servizi-sociali-per-abitante.xlsx"
BENCHMARK_URL="https://www.istat.it/wp-content/uploads/2025/09/Tavole_2022_spesa_comuni-1.xlsx"
LANDING="https://www.istat.it/comunicato-stampa/la-spesa-dei-comuni-per-i-servizi-sociali-anno-2022/"
METRIC="socialSpendingPerResident"
BLOCKED_METRIC="socialSpendingByUserArea"

def finite(v:Any)->float:
    if isinstance(v,bool):
        raise ValueError(v)
    x=float(v)
    if not math.isfinite(x):
        raise ValueError(v)
    return x

def fetch_book(session:requests.Session,url:str):
    r=session.get(url,timeout=240)
    r.raise_for_status()
    return load_workbook(io.BytesIO(r.content),read_only=True,data_only=True),len(r.content)

def reconcile_municipal_source(wb, public_rows:list[dict])->dict:
    evidence={}
    errors=[]
    for item in public_rows:
        town=str(item.get("town") or "").strip()
        expected=finite(item.get("value"))
        candidates=[]
        target=town.casefold()
        for ws in wb.worksheets:
            for idx,row in enumerate(ws.iter_rows(values_only=True),start=1):
                cells=list(row)
                if not any(str(v or "").strip().casefold()==target for v in cells):
                    continue
                nums=[]
                for value in cells:
                    if isinstance(value,(int,float)) and not isinstance(value,bool) and math.isfinite(float(value)):
                        nums.append(float(value))
                matches=[v for v in nums if math.isclose(v,expected,rel_tol=0.0,abs_tol=.011)]
                if matches:
                    candidates.append({"sheet":ws.title,"row":idx,"value":matches[0]})
        distinct={(x["sheet"],x["row"],round(x["value"],8)) for x in candidates}
        if len(distinct)!=1:
            errors.append(f"{town}: riconciliazione fonte comunale non univoca {sorted(distinct)}")
        else:
            evidence[town]=next(iter(candidates))
    if errors:
        raise RuntimeError("; ".join(errors[:20]))
    if len(evidence)!=7:
        raise RuntimeError(f"fonte comunale: riconciliazione {len(evidence)}/7")
    return evidence

def benchmark_rows(wb)->tuple[dict,dict]:
    if "Tav. 1" not in wb.sheetnames:
        raise RuntimeError(f"Istat sociale: Tav. 1 assente {wb.sheetnames}")
    ws=wb["Tav. 1"]
    found={}
    rows={}
    for idx,row in enumerate(ws.iter_rows(values_only=True),start=1):
        cells=list(row)
        label=str(cells[0] or "").strip().casefold() if cells else ""
        key=None
        if label=="toscana":
            key="tuscany"
        elif label=="italia":
            key="italy"
        if key:
            # Tav. 1: colonna F = Spesa pro-capite, inclusi i servizi educativi per la prima infanzia.
            if len(cells)<6:
                raise RuntimeError(f"Istat sociale {key}: riga corta")
            value=finite(cells[5])
            found[key]=value
            rows[key]={"sheet":ws.title,"row":idx,"values":cells[:8]}
    if set(found)!={"tuscany","italy"}:
        raise RuntimeError(f"Istat sociale: righe Toscana/Italia incomplete {found}")
    return found,rows

def main()->None:
    ap=argparse.ArgumentParser()
    ap.add_argument("--output",required=True)
    args=ap.parse_args()

    site=json.loads(SITE.read_text(encoding="utf-8"))
    metric=(site.get("metrics") or {}).get(METRIC) or {}
    meta=metric.get("meta") or {}
    public_rows=metric.get("rows") or []
    if len(public_rows)!=7:
        raise RuntimeError(f"{METRIC}: pubblico {len(public_rows)}/7")
    if str(meta.get("unit") or "")!="eurPerResident" or str(meta.get("year") or "")!="2022":
        raise RuntimeError(f"{METRIC}: contratto pubblico inatteso {meta.get('unit')}/{meta.get('year')}")
    formula=str((metric.get("method") or {}).get("formula") or "").casefold()
    if "spesa dei comuni" not in formula or "popolazione residente media" not in formula:
        raise RuntimeError(f"{METRIC}: formula pubblica inattesa {formula}")
    caveat=str((metric.get("method") or {}).get("caveat") or "").casefold()
    if "include i servizi educativi per la prima infanzia" not in caveat:
        raise RuntimeError(f"{METRIC}: perimetro servizi educativi non certificato")

    session=requests.Session()
    session.headers["User-Agent"]="OsservatorioVersilia-A3-social-benchmark/1.0"
    municipal_wb,municipal_bytes=fetch_book(session,MUNICIPAL_URL)
    municipal_evidence=reconcile_municipal_source(municipal_wb,public_rows)

    benchmark_wb,benchmark_bytes=fetch_book(session,BENCHMARK_URL)
    scope_values,scope_rows=benchmark_rows(benchmark_wb)
    if not (0<scope_values["tuscany"]<1000 and 0<scope_values["italy"]<1000):
        raise RuntimeError(f"Istat sociale: benchmark fuori intervallo {scope_values}")

    payload={
      "schemaVersion":1,
      "publisher":"Istat — Spesa dei Comuni per i servizi sociali",
      "profileId":"istat-social-services-annual",
      "referenceYear":2022,
      "status":"ACQUIRED_CANDIDATE",
      "sourceUrl":LANDING,
      "municipalDataUrl":MUNICIPAL_URL,
      "benchmarkDataUrl":BENCHMARK_URL,
      "benchmarks":{
        METRIC:{
          "year":"2022",
          "unit":"eurPerResident",
          "formula":"spesa per interventi e servizi sociali dei comuni singoli e associati / popolazione residente media",
          "tuscany":scope_values["tuscany"],
          "italy":scope_values["italy"],
        }
      },
      "qualityGate":{
        "status":"PASS",
        "publicReconciliation":"1 metric × 7/7 towns PASS against Istat A misura di Comune",
        "officialBenchmarkRows":"Toscana + Italia PASS from Tav. 1, municipal social expenditure only",
        "definitionControl":"includes early-childhood educational services, matching public metric caveat",
        "errors":[],
      },
      "evidence":{
        "municipalRows":municipal_evidence,
        "benchmarkRows":scope_rows,
        "municipalBytes":municipal_bytes,
        "benchmarkBytes":benchmark_bytes,
      },
      "blocked":{
        BLOCKED_METRIC:"metrica composita a 7 aree di utenza; meta.benchmark scalare unico non è compatibile con il selettore per componente"
      }
    }
    out=Path(args.output)
    out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({
      "status":payload["status"],
      "candidateMetrics":[METRIC],
      "tuscany":scope_values["tuscany"],
      "italy":scope_values["italy"],
      "blocked":payload["blocked"],
    },ensure_ascii=False))

if __name__=="__main__":
    main()
