#!/usr/bin/env python3
from __future__ import annotations

import argparse, hashlib, io, json, math, zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

ROOT=Path(__file__).resolve().parents[1]
SITE=ROOT/"data/site-data.json"
DEMO=ROOT/"data/source-snapshots/a3-istat-demography-benchmark-2026.json"
URLS={
 "movement":"https://www.regione.toscana.it/documents/d/guest/2-movimento-per-comune-2025-agg-maggio-2026-",
 "monthly":"https://www.regione.toscana.it/documents/d/guest/5-movimento-comune_mese-2025-agg-maggio-2026-",
 "capacity":"https://www.regione.toscana.it/documents/d/guest/1-consistenza-media-per-comune-e-tipologia-ricettiva-2025",
 "rentals":"https://www.regione.toscana.it/documents/d/guest/6-locazioni-agg-maggio-2026-",
}
NS={"office":"urn:oasis:names:tc:opendocument:xmlns:office:1.0","table":"urn:oasis:names:tc:opendocument:xmlns:table:1.0","text":"urn:oasis:names:tc:opendocument:xmlns:text:1.0"}
TARGETS=["foreignTourismShare","tourismArrivals","tourismAverageStay","tourismIntensity","tourismPresences","tourismSeasonality"]

def cell_value(cell:ET.Element)->Any:
    typ=cell.attrib.get(f"{{{NS['office']}}}value-type")
    if typ in {"float","currency","percentage"}:
        raw=cell.attrib.get(f"{{{NS['office']}}}value")
        if raw is not None:
            try:return float(raw)
            except ValueError: pass
    texts=[]
    for p in cell.findall(".//text:p",NS): texts.append("".join(p.itertext()))
    return " ".join(x for x in texts if x).strip()

def ods_rows(blob:bytes)->list[list[Any]]:
    with zipfile.ZipFile(io.BytesIO(blob)) as z:
        root=ET.fromstring(z.read("content.xml"))
    tables=root.findall(".//table:table",NS)
    if len(tables)!=1: raise RuntimeError(f"ODS: atteso un foglio, trovati {len(tables)}")
    rows=[]
    for tr in tables[0].findall("table:table-row",NS):
        repeat=int(tr.attrib.get(f"{{{NS['table']}}}number-rows-repeated","1"))
        vals=[]
        for cell in list(tr):
            if cell.tag not in {f"{{{NS['table']}}}table-cell",f"{{{NS['table']}}}covered-table-cell"}: continue
            rep=int(cell.attrib.get(f"{{{NS['table']}}}number-columns-repeated","1"))
            vals.extend([cell_value(cell)]*min(rep,60))
        while vals and vals[-1] in ("",None): vals.pop()
        if vals:
            rows.extend([vals]*min(repeat,2))
    return rows

def fetch(url:str)->tuple[list[list[Any]],int]:
    import requests
    r=requests.get(url,timeout=180,headers={"User-Agent":"OsservatorioVersilia-A3-tourism/2.0"}); r.raise_for_status()
    return ods_rows(r.content),len(r.content)

def find_rows(rows:list[list[Any]])->dict[str,list[Any]]:
    out={}
    for row in rows:
        if not row: continue
        name=str(row[1] if len(row)>1 and row[1] else row[0]).strip()
        if name:
            out[name.casefold()]=row
    return out

def number(v:Any,label:str)->float:
    if isinstance(v,bool): raise RuntimeError(f"{label}: boolean")
    try:x=float(v)
    except Exception as exc: raise RuntimeError(f"{label}: numerico atteso {v!r}") from exc
    if not math.isfinite(x): raise RuntimeError(f"{label}: non finito")
    return x

def town_public(site:dict,metric_id:str)->dict[str,float]:
    metric=site["metrics"][metric_id]
    rows={str(r["town"]):float(r["value"]) for r in metric.get("rows",[]) if r.get("town") and isinstance(r.get("value"),(int,float))}
    if len(rows)!=7: raise RuntimeError(f"{metric_id}: pubblico non 7/7")
    return rows

def rows_named(rows:list[list[Any]],name:str)->list[list[Any]]:
    target=name.strip().casefold()
    out=[]
    for row in rows:
        if any(str(cell or "").strip().casefold()==target for cell in row):
            out.append(row)
    return out

def capacity_component(
    rows:list[list[Any]],
    metric_id:str,
    site:dict,
    pop_by_town:dict[str,float],
    town_names:list[str],
)->dict[str,Any]:
    public=town_public(site,metric_id)
    hits={}
    for town in town_names:
        candidates=[]
        for row in rows_named(rows,town):
            for col,value in enumerate(row):
                if isinstance(value,bool): continue
                try: absolute=float(value)
                except Exception: continue
                if absolute<0: continue
                rate=absolute/pop_by_town[town]*1000.0
                if math.isclose(rate,public[town],rel_tol=0.0,abs_tol=.11):
                    candidates.append((col,absolute,row))
        hits[town]=candidates
    common=None
    for town,candidates in hits.items():
        cols={col for col,_,_ in candidates}
        common=cols if common is None else common & cols
    if not common:
        raise RuntimeError(
            f"{metric_id}: nessuna colonna ODS comune riconcilia 7/7; "
            + " | ".join(f"{town}:{[(c,v) for c,v,_ in cand[:8]]}" for town,cand in hits.items())
        )
    resolved=[]
    for col in sorted(common):
        values={}
        ok=True
        for town,candidates in hits.items():
            nums=sorted({float(v) for c,v,_ in candidates if c==col})
            if len(nums)!=1:
                ok=False; break
            values[town]=nums[0]
        if ok: resolved.append((col,values))
    if len(resolved)!=1:
        raise RuntimeError(f"{metric_id}: colonne riconciliate non univoche {[(c,v) for c,v in resolved]}")
    col,values=resolved[0]

    regional=[]
    for row in rows_named(rows,"Toscana"):
        if col>=len(row): continue
        try: value=float(row[col])
        except Exception: continue
        labels=" ".join(str(x or "").strip().casefold() for x in row)
        regional.append((value,labels,row))
    total_rows=[item for item in regional if "totale" in item[1]]
    candidates=total_rows or regional
    distinct=sorted({float(value) for value,_,_ in candidates if value>=0})
    if len(distinct)!=1:
        raise RuntimeError(f"{metric_id}: aggregato Toscana non univoco col={col}, values={distinct[:30]}")
    return {"column":col,"townAbsolute":values,"tuscanyAbsolute":distinct[0]}

HISTORY_SNAPSHOT = "data/source-snapshots/a3-regione-toscana-tourism-benchmark-2025.json"
HISTORY_URLS = {
    2023: "https://www.regione.toscana.it/documents/10180/11976751/Movimento%2Bper%2Bcomune%2B2023.xlsx/68162001-ccd5-6879-81c9-99639ee947bb?t=1709731593776",
    2024: "https://www.regione.toscana.it/documents/d/guest/2-movimento-per-comune-2024-ods",
    2025: URLS["movement"],
}
HISTORY_METRICS = ("tourismArrivals", "tourismPresences", "tourismAverageStay", "foreignTourismShare")


def parse_movement_history(bodies: dict[int, bytes], towns: dict[str, str]) -> dict:
    """Read three small municipal tables; preserve the native counts and provenance."""
    if set(bodies) != set(HISTORY_URLS):
        raise RuntimeError("Movement history source periods mismatch")
    annual = {}
    sources = {}
    for year, body in sorted(bodies.items()):
        if year == 2023:
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
            try:
                rows = [list(row) for row in book.active.iter_rows(values_only=True)]
            finally:
                book.close()
        else:
            rows = ods_rows(body)
        header = " ".join(str(cell or "") for row in rows[:6] for cell in row).casefold()
        if str(year) not in header or "al netto delle locazioni" not in header:
            raise RuntimeError("Movement history source scope/year mismatch")
        annual[str(year)] = {}
        for code, town in towns.items():
            hits = [row for row in rows if len(row) >= 8 and str(row[1] or "").strip().casefold() == town.casefold()]
            if len(hits) != 1:
                raise RuntimeError(f"Movement history municipality not unique: {year}/{town}")
            counts = []
            for value in hits[0][2:8]:
                x = number(value, f"{year}/{town}")
                if x < 0 or not x.is_integer():
                    raise RuntimeError("Movement history count invalid")
                counts.append(int(x))
            ai, af, arrivals, pi, pf, presences = counts
            if ai + af != arrivals or pi + pf != presences or arrivals <= 0 or presences <= 0:
                raise RuntimeError("Movement history component totals mismatch")
            annual[str(year)][code] = {"arrivals": arrivals, "presences": presences, "foreignPresences": pf}
        sources[str(year)] = {"url": HISTORY_URLS[year], "sha256": hashlib.sha256(body).hexdigest(), "bytes": len(body)}
    return {"years": sorted(bodies), "scope": "al netto delle locazioni", "sources": sources, "countsByYear": annual}


def apply_movement_history(metric: dict, snapshot: dict) -> None:
    history = snapshot.get("municipalMovementHistory")
    if history is None:
        return
    mid = metric["meta"]["key"]
    if mid not in HISTORY_METRICS or str(metric["meta"]["year"]) != "2025":
        raise RuntimeError("Movement history metric/year mismatch")
    years = history.get("years")
    if years != [2023, 2024, 2025] or history.get("scope") != "al netto delle locazioni":
        raise RuntimeError("Movement history period/scope mismatch")
    rows = metric["rows"]
    codes = {r["code"] for r in rows}
    if len(rows) != 7 or len(codes) != 7 or set(history["countsByYear"]) != {str(y) for y in years}:
        raise RuntimeError("Movement history public scope mismatch")
    values = {code: [] for code in codes}
    aggregate_values = []
    for year in years:
        provenance = history["sources"][str(year)]
        if provenance["url"] != HISTORY_URLS[year] or len(provenance["sha256"]) != 64 or any(c not in "0123456789abcdef" for c in provenance["sha256"]) or provenance["bytes"] <= 0:
            raise RuntimeError("Movement history provenance invalid")
        counts = history["countsByYear"][str(year)]
        if set(counts) != codes:
            raise RuntimeError("Movement history annual coverage mismatch")
        for code, components in counts.items():
            if set(components) != {"arrivals", "presences", "foreignPresences"}:
                raise RuntimeError("Movement history components missing")
            a, p, f = (components[k] for k in ("arrivals", "presences", "foreignPresences"))
            if any(isinstance(v, bool) or not isinstance(v, int) or v < 0 for v in (a, p, f)) or a <= 0 or p <= 0 or f > p:
                raise RuntimeError("Movement history components invalid")
            values[code].append({"tourismArrivals": a, "tourismPresences": p,
                                 "tourismAverageStay": p / a, "foreignTourismShare": round(f / p * 100, 1)}[mid])
        total_a = sum(c["arrivals"] for c in counts.values())
        total_p = sum(c["presences"] for c in counts.values())
        # The current public foreign share weights municipal percentages rounded to one decimal.
        aggregate_values.append({"tourismArrivals": total_a, "tourismPresences": total_p,
                                 "tourismAverageStay": total_p / total_a,
                                 "foreignTourismShare": sum(values[code][-1] * counts[code]["presences"] for code in codes) / total_p}[mid])
    metadata = {"source": snapshot["publisher"], "sourceUrl": snapshot["sourceUrl"],
                "sourceSnapshot": HISTORY_SNAPSHOT,
                "note": "Movimento turistico 2023–2025 al netto delle locazioni; tavole regionali provvisorie fino alla diffusione Istat. Quota estera comunale arrotondata a un decimale; quota Versilia ponderata sulle presenze con lo stesso criterio del valore corrente."}
    candidates = [(row, values[row["code"]]) for row in rows] + [(metric["aggregate"], aggregate_values)]
    for holder, observations in candidates:
        if not math.isclose(observations[-1], holder["value"], rel_tol=0, abs_tol=1e-9):
            raise RuntimeError("Movement history latest public value mismatch")
        existing = holder.get("series") or {}
        old_years, old_values = existing.get("years", []), existing.get("values", [])
        if len(old_years) != len(old_values) or any(y not in years or not math.isclose(v, observations[years.index(y)], rel_tol=0, abs_tol=1e-9) for y, v in zip(old_years, old_values)):
            raise RuntimeError("Movement history conflicts with existing public series")
    for holder, observations in candidates:
        holder["series"] = {"years": years, "values": observations, **metadata}


def acquire_movement_history(site: dict) -> dict:
    import requests
    bodies = {}
    for year, url in HISTORY_URLS.items():
        response = requests.get(url, timeout=45)
        response.raise_for_status()
        if len(response.content) > 2 * 1024 * 1024:
            raise RuntimeError("Movement history source exceeds simple-table size limit")
        bodies[year] = response.content
    towns = {row["code"]: row["town"] for row in site["metrics"]["tourismArrivals"]["rows"]}
    return parse_movement_history(bodies, towns)


def main()->None:
    ap=argparse.ArgumentParser(); ap.add_argument("--output",required=True); ap.add_argument("--history-only", action="store_true"); a=ap.parse_args()
    site=json.loads(SITE.read_text(encoding="utf-8"))
    if a.history_only:
        payload = json.loads((ROOT / HISTORY_SNAPSHOT).read_text(encoding="utf-8"))
        payload["municipalMovementHistory"] = acquire_movement_history(site)
        for mid in HISTORY_METRICS:
            apply_movement_history(site["metrics"][mid], payload)
        Path(a.output).write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print("Movement history PASS: 2023–2025 × 7 municipalities; four metrics reconciled")
        return
    demo=json.loads(DEMO.read_text(encoding="utf-8"))
    movement,mb=fetch(URLS["movement"]); monthly,monb=fetch(URLS["monthly"]); capacity,capb=fetch(URLS["capacity"]); rentals,renb=fetch(URLS["rentals"])
    mi=find_rows(movement); mo=find_rows(monthly)
    town_names=[str(t["name"]) for t in site.get("towns",[])]
    pop_by_town={str(r["town"]):float(r["value"]) for r in site["metrics"]["population"]["rows"]}
    errors=[]
    computed={mid:{} for mid in TARGETS}
    for town in town_names:
        mr=mi.get(town.casefold()); xr=mo.get(town.casefold())
        if mr is None or xr is None:
            errors.append(f"{town}: riga movement/monthly mancante"); continue
        arrivals=number(mr[4],f"{town}/arrivals")
        foreign_pres=number(mr[6],f"{town}/foreign-presences")
        presences=number(mr[7],f"{town}/presences")
        month_pres=[number(xr[i],f"{town}/month-{m}") for m,i in enumerate(range(3,26,2),1)]
        monthly_total=number(xr[27],f"{town}/monthly-total")
        if not math.isclose(sum(month_pres),monthly_total,rel_tol=0.0,abs_tol=.1) or not math.isclose(monthly_total,presences,rel_tol=0.0,abs_tol=.1):
            errors.append(f"{town}: mensile non riconciliato")
        computed["tourismArrivals"][town]=arrivals
        computed["tourismPresences"][town]=presences
        computed["tourismAverageStay"][town]=presences/arrivals
        computed["foreignTourismShare"][town]=foreign_pres/presences*100
        computed["tourismSeasonality"][town]=sum(sorted(month_pres,reverse=True)[:3])/presences*100
        computed["tourismIntensity"][town]=presences/pop_by_town[town]

    tolerances={"tourismArrivals":.1,"tourismPresences":.1,"tourismAverageStay":1e-9,"foreignTourismShare":.11,"tourismSeasonality":1e-9,"tourismIntensity":1e-9}
    for mid in TARGETS:
        public=town_public(site,mid)
        for town,expected in computed[mid].items():
            if not math.isclose(public[town],expected,rel_tol=0.0,abs_tol=tolerances[mid]):
                errors.append(f"{mid}/{town}: {public[town]} != {expected}")

    tr=mi.get("toscana")
    xr=mo.get("toscana")
    if tr is None or xr is None: errors.append("Toscana: riga aggregata mancante")
    benchmarks={}
    if tr is not None and xr is not None:
        arrivals=number(tr[4],"Toscana/arrivals"); foreign_pres=number(tr[6],"Toscana/foreign-presences"); pres=number(tr[7],"Toscana/presences")
        month_pres=[number(xr[i],f"Toscana/month-{m}") for m,i in enumerate(range(3,26,2),1)]
        if not math.isclose(sum(month_pres),pres,rel_tol=0.0,abs_tol=.1): errors.append("Toscana: mensile != presenze annuali")
        pop=float(demo["benchmarks"]["population"]["tuscany"])
        benchmarks={
          "tourismArrivals":{"year":"2025","unit":"number","tuscany":arrivals,"italy":None},
          "tourismPresences":{"year":"2025","unit":"number","tuscany":pres,"italy":None},
          "tourismAverageStay":{"year":"2025","unit":"nights","formula":"presenze / arrivi","tuscany":pres/arrivals,"italy":None},
          "foreignTourismShare":{"year":"2025","unit":"percent","formula":"presenze straniere / presenze totali × 100","tuscany":foreign_pres/pres*100,"italy":None},
          "tourismSeasonality":{"year":"2025","unit":"percent","formula":"tre mesi con più presenze / presenze annue × 100","tuscany":sum(sorted(month_pres,reverse=True)[:3])/pres*100,"italy":None},
          "tourismIntensity":{"year":"2025","unit":"decimal","formula":"presenze 2025 / popolazione benchmark 2026","tuscany":pres/pop,"italy":None},
        }

    movement_errors=list(errors)
    blocked={
      "italy":"la stessa fonte Regione Toscana non espone un aggregato nazionale omogeneo; benchmark Italia resta n.d."
    }
    capacity_gate={}
    capacity_diagnostics={
      name:[[cell for cell in row[:40]] for row in rows_named(capacity,name)[:12]]
      for name in [*town_names,"Toscana"]
    }
    rental_diagnostics={
      name:[[cell for cell in row[:40]] for row in rows_named(rentals,name)[:20]]
      for name in [*town_names,"Toscana"]
    }
    rental_preview=[
      [cell for cell in row[:40]]
      for row in rentals[:80]
    ]
    for metric_id in ("tourismBedsPer1000","tourismStructuresPer1000"):
        try:
            detail=capacity_component(capacity,metric_id,site,pop_by_town,town_names)
            pop=float(demo["benchmarks"]["population"]["tuscany"])
            benchmarks[metric_id]={
              "year":"2025","unit":"per1000",
              "formula":"consistenza ricettiva regionale / popolazione benchmark 2026 × 1.000",
              "tuscany":detail["tuscanyAbsolute"]/pop*1000.0,"italy":None,
            }
            capacity_gate[metric_id]={"status":"PASS",**detail}
        except Exception as exc:
            capacity_gate[metric_id]={"status":"BLOCKED","reason":str(exc)}
            blocked[metric_id]=str(exc)
    gate="PASS" if not movement_errors else "FAIL"
    payload={
      "schemaVersion":3,"publisher":"Regione Toscana — Ufficio regionale di Statistica / Istat","profileId":"regione-toscana-tourism-annual","referenceYear":2025,
      "status":"ACQUIRED_CANDIDATE" if gate=="PASS" else "CANDIDATE_REJECTED",
      "sources":{"movement":URLS["movement"],"monthly":URLS["monthly"],"capacity":URLS["capacity"],"rentals":URLS["rentals"],"movementBytes":mb,"monthlyBytes":monb,"capacityBytes":capb,"rentalsBytes":renb},
      "benchmarks":benchmarks,
      "qualityGate":{"status":gate,"publicReconciliation":"6 movement metrics × 7/7 towns PASS" if gate=="PASS" else "FAIL","regionalRows":"movement + monthly Toscana PASS" if gate=="PASS" else "FAIL","capacityMetricGates":capacity_gate,"errors":movement_errors},
      "capacityDiagnostics":capacity_diagnostics,
      "rentalDiagnostics":rental_diagnostics,
      "rentalPreview":rental_preview,
      "blocked":blocked
    }
    p=Path(a.output); p.parent.mkdir(parents=True,exist_ok=True); p.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"status":payload["status"],"candidateMetrics":sorted(benchmarks),"tuscany":{k:v["tuscany"] for k,v in benchmarks.items()},"gate":payload["qualityGate"]},ensure_ascii=False))
if __name__=="__main__": main()
