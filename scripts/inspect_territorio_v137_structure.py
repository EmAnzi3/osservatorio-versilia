#!/usr/bin/env python3
"""Ispeziona struttura e attributi delle fonti ufficiali v1.37.

Produce un JSON diagnostico, senza modificare i dati pubblicati. L'obiettivo e'
bloccare nomi campo, CRS, filtri e contenuto degli archivi prima del generatore GIS.
"""
from __future__ import annotations

import io
import json
import re
import shutil
import tempfile
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path

import geopandas as gpd
from openpyxl import load_workbook
import pyogrio

OUT = Path("tmp/territorio-v137-structure.json")
TOWNS = ["Camaiore", "Forte dei Marmi", "Massarosa", "Pietrasanta", "Seravezza", "Stazzema", "Viareggio"]
UA = {"User-Agent": "OsservatorioVersilia-v137-structure-audit/1.0"}
ISPRA = "https://www.isprambiente.gov.it/it/attivita/suolo-e-territorio/suolo/il-consumo-di-suolo/consumo_di_suolo_estratto_dati_2025_anni_2006_2024.xlsx"
RETICULUM = "https://www.regione.toscana.it/documents/d/guest/infrastruttura_rev25-zip"
ROADS = "https://www502.regione.toscana.it/geoscopio/download/grafo_stradale/iternet.zip"
WFS_BASE = "https://www502.regione.toscana.it/wmsraster/com.rt.wms.RTmap/wms"
WFS_LAYERS = {
    "parksNational": "rt_arprot.idparnaz.rt.poly",
    "reservesNational": "rt_arprot.idrisnatstat.rt.poly",
    "parksRegional": "rt_arprot.idparreg.rt.poly",
    "parksProvincial": "rt_arprot.idparprov.rt.poly",
    "reservesRegional": "rt_arprot.idrisnatreg.rt.poly",
    "anpil": "rt_arprot.idanpil.rt.poly",
    "zsc": "rt_arprot.idnat2000_sic.rt.poly",
    "zps": "rt_arprot.idnat2000_zps.rt.poly",
    "zscZps": "rt_arprot.idnat2000_sic_zps.rt.poly",
    "ramsar": "rt_arprot.idramsar.rt.poly",
}


def download(url: str, dest: Path) -> None:
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=600) as response, dest.open("wb") as out:
        shutil.copyfileobj(response, out, length=1024 * 1024)


def wfs_url(layer: str) -> str:
    params = {
        "map": "wmsarprot", "service": "WFS", "version": "2.0.0", "request": "GetFeature",
        "OUTPUTFORMAT": "application/shapefile", "filename": "v137_structure", "typename": layer,
    }
    return WFS_BASE + "?" + urllib.parse.urlencode(params)


def json_value(value):
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def ispra_structure(raw: bytes) -> dict:
    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb["Comuni_2006_2024"]
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    selected = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[1] in TOWNS and row[3] == "Lucca":
            selected.append({str(headers[i]): json_value(row[i]) for i in range(len(headers))})
    if [row["Nome_Comune"] for row in selected] != TOWNS:
        selected = sorted(selected, key=lambda row: TOWNS.index(row["Nome_Comune"]))
    return {"headers": headers, "selectedRows": selected, "selectedCount": len(selected)}


def vector_info(path: str) -> dict:
    info = pyogrio.read_info(path)
    return {
        "crs": str(info.get("crs")),
        "featureCount": int(info.get("features", 0)),
        "geometryType": str(info.get("geometry_type")),
        "fields": [str(x) for x in info.get("fields", [])],
        "dtypes": [str(x) for x in info.get("dtypes", [])],
    }


def unique_preview(frame: gpd.GeoDataFrame, field: str, limit: int = 100) -> list:
    if field not in frame.columns:
        return []
    values = frame[field].dropna().astype(str).value_counts(dropna=False)
    return [{"value": index, "count": int(count)} for index, count in values.head(limit).items()]


def reticulum_structure(zip_path: Path, work: Path) -> dict:
    target = work / "reticulum"
    target.mkdir()
    with zipfile.ZipFile(zip_path) as archive:
        archive.extractall(target)
    shp = target / "reticoloDCR242025.shp"
    info = vector_info(str(shp))
    frame = gpd.read_file(shp, engine="pyogrio")
    info["COMPLR79"] = unique_preview(frame, "COMPLR79")
    info["RETGESLR79"] = unique_preview(frame, "RETGESLR79")
    info["managedFilterCount"] = int(((frame.get("COMPLR79").astype(str) == "Toscana Nord") & (frame.get("RETGESLR79").astype(str) == "SI")).sum()) if {"COMPLR79", "RETGESLR79"}.issubset(frame.columns) else None
    return info


def protected_structure(work: Path) -> dict:
    result = {}
    for key, layer in WFS_LAYERS.items():
        raw = urllib.request.urlopen(urllib.request.Request(wfs_url(layer), headers=UA), timeout=240).read()
        folder = work / f"protected-{key}"
        folder.mkdir()
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            archive.extractall(folder)
        shp = next(folder.glob("*.shp"))
        info = vector_info(str(shp))
        frame = gpd.read_file(shp, engine="pyogrio")
        for field in ("STATO", "TIPO", "DENOMINAZ", "DENOM", "NOME", "CODICE"):
            if field in frame.columns:
                info[f"unique:{field}"] = unique_preview(frame, field, 40)
        result[key] = info
    return result


def roads_structure(outer_zip: Path, work: Path) -> dict:
    inner_path = work / "iternet-inner.zip"
    with zipfile.ZipFile(outer_zip) as archive:
        members = archive.namelist()
        inner = next((name for name in members if name.lower().endswith(".zip")), None)
        if not inner:
            raise RuntimeError("Iter.Net: archivio interno .zip non trovato")
        with archive.open(inner) as src, inner_path.open("wb") as dst:
            shutil.copyfileobj(src, dst, length=1024 * 1024)
    with zipfile.ZipFile(inner_path) as archive:
        shp_names = [name for name in archive.namelist() if name.lower().endswith(".shp")]
    likely = [name for name in shp_names if re.search(r"(elem|el_?str|strad|arco)", Path(name).stem, re.I)]
    infos = {}
    for name in likely[:30]:
        try:
            infos[name] = vector_info(f"/vsizip/{inner_path.resolve()}/{name}")
        except Exception as exc:  # diagnostics: preserve exact failure
            infos[name] = {"error": repr(exc)}
    return {
        "outerBytes": outer_zip.stat().st_size,
        "innerBytes": inner_path.stat().st_size,
        "shapefiles": shp_names,
        "likelyRoadLayers": likely,
        "likelyRoadLayerInfo": infos,
    }


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="territorio-v137-structure-") as tmp:
        work = Path(tmp)
        ispra_path = work / "ispra.xlsx"
        reticulum_path = work / "reticulum.zip"
        roads_path = work / "iternet.zip"

        download(ISPRA, ispra_path)
        download(RETICULUM, reticulum_path)
        download(ROADS, roads_path)

        payload = {
            "ispra": ispra_structure(ispra_path.read_bytes()),
            "reticulum": reticulum_structure(reticulum_path, work),
            "protected": protected_structure(work),
            "roads": roads_structure(roads_path, work),
        }

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))

    if payload["ispra"]["selectedCount"] != 7:
        raise SystemExit(f"ISPRA: trovati {payload['ispra']['selectedCount']} Comuni su 7")
    if not payload["roads"]["likelyRoadLayers"]:
        raise SystemExit("Iter.Net: nessun layer stradale candidato identificato")


if __name__ == "__main__":
    main()
