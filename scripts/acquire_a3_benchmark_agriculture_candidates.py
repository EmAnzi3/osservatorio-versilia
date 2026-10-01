#!/usr/bin/env python3
from __future__ import annotations
import argparse, io, json, re
from pathlib import Path
from typing import Any
import requests
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
LOCAL=ROOT/'data/source-snapshots/istat-agricoltura-territorio-2020.json'
WORKBOOK_URL='https://www.regione.toscana.it/documents/10180/12003883/TOSCANA_AGRICOLTURA_CENSIMENTO2020.xlsx/8e641f76-e024-4d15-b722-5afb9aec5c8e?t=1708419413979'
SDMX='https://esploradati.istat.it/SDMXWS/rest'
TOWN_QUERY='046005+046013+046018+046024+046028+046030+046033'

def norm(v:Any)->str:
    s=str(v or '').strip().casefold()
    s=re.sub(r'\s+',' ',s)
    return s

def serial(v:Any)->Any:
    if v is None or isinstance(v,(str,int,float,bool)): return v
    return str(v)

def row_values(ws,rownum:int,maxcol:int)->list[Any]:
    return [serial(ws.cell(rownum,c).value) for c in range(1,maxcol+1)]

def workbook_scan(blob:bytes)->dict[str,Any]:
    wb=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    try:
        sheets=[]
        for ws in wb.worksheets:
            maxcol=min(ws.max_column or 1,90)
            maxrow=min(ws.max_row or 1,5000)
            top=[]
            for r in range(1,min(maxrow,18)+1):
                vals=row_values(ws,r,maxcol)
                text=' | '.join(str(x) for x in vals if x not in (None,'')).strip()
                if text: top.append(text[:1200])
            matches=[]
            for r in range(1,maxrow+1):
                vals=row_values(ws,r,maxcol)
                normalized=[norm(x) for x in vals]
                labels=[x for x in normalized if x]
                target=None
                if any(x=='toscana' for x in labels): target='Toscana'
                elif any(x in {'italia','totale italia'} for x in labels): target='Italia'
                if not target: continue
                window=[]
                for rr in range(max(1,r-8),r+1):
                    window.append({'row':rr,'values':row_values(ws,rr,maxcol)})
                matches.append({'target':target,'row':r,'values':vals,'headerWindow':window})
            if matches:
                sheets.append({
                    'title':ws.title,
                    'maxRow':ws.max_row,
                    'maxColumn':ws.max_column,
                    'topText':top,
                    'matches':matches[:20],
                })
        return {'sheetNames':wb.sheetnames,'matchedSheets':sheets}
    finally:
        wb.close()

def sdmx_probe(session:requests.Session,flow:str,key:str)->dict[str,Any]:
    url=f'{SDMX}/data/{flow}/{key}/IT1'
    try:
        r=session.get(url,params={'startPeriod':'2020','endPeriod':'2020','format':'csvfile'},timeout=180)
        r.raise_for_status()
        text=r.content.decode('utf-8-sig',errors='replace')
        lines=text.splitlines()
        return {'url':r.url,'status':r.status_code,'bytes':len(r.content),'preview':lines[:12]}
    except Exception as exc:
        return {'url':url,'error':f'{type(exc).__name__}: {exc}'}

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--output',required=True); a=ap.parse_args()
    session=requests.Session()
    session.headers.update({'User-Agent':'OsservatorioVersilia-A3-agriculture/2.0'})
    response=session.get(WORKBOOK_URL,timeout=240)
    response.raise_for_status()
    scan=workbook_scan(response.content)
    local=json.loads(LOCAL.read_text(encoding='utf-8'))
    probes={
      'surface7':sdmx_probe(session,'DF_DCAT_CENSAGRIC2020_SURF_ALL',f'A.{TOWN_QUERY}.HO+ARU+FUAA'),
      'irrigation7':sdmx_probe(session,'DF_DCAT_CENSAGRIC2020_SURF_IRR_CONS',f'A.{TOWN_QUERY}.IA'),
      'localized7':sdmx_probe(session,'DF_DCAT_CENSAGRIC2020_UA_CROPS_2',f'A.{TOWN_QUERY}.ARU.ALL.TOT'),
    }
    payload={
      'schemaVersion':2,
      'publisher':'Istat / Regione Toscana — Ufficio di Statistica',
      'profileId':'istat-agriculture-census-2020',
      'referenceYear':2020,
      'status':'SOURCE_DIAGNOSTIC_READY',
      'source':{'url':WORKBOOK_URL,'bytes':len(response.content)},
      'workbook':scan,
      'sdmxProbes':probes,
      'publicSnapshot':{
        'coverage':local.get('verification',{}).get('result'),
        'townCount':len(local.get('towns') or {}),
        'definitions':local.get('definitions'),
        'derivations':local.get('derivations'),
      },
      'goal':'Identificare nello stesso artifact le righe Toscana/Italia delle tavole ufficiali e la forma reale delle risposte SDMX 7/7; nessun benchmark viene pubblicato finché definizione e denominatore non coincidono con il contratto pubblico.',
    }
    p=Path(a.output); p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'status':payload['status'],'matchedSheets':len(scan['matchedSheets']),'sheetNames':scan['sheetNames']},ensure_ascii=False))
if __name__=='__main__': main()
