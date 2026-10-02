#!/usr/bin/env python3
from __future__ import annotations

import argparse,io,json,math,re,unicodedata
from pathlib import Path
from typing import Any
import requests
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
SITE=ROOT/"data/site-data.json"
URL="https://www.isprambiente.gov.it/it/attivita/suolo-e-territorio/suolo/il-consumo-di-suolo/consumo_di_suolo_estratto_dati_2025_anni_2006_2024.xlsx"
LANDING="https://www.isprambiente.gov.it/it/attivita/suolo-e-territorio/suolo/il-consumo-di-suolo/i-dati-sul-consumo-di-suolo"

def norm(v:Any)->str:
    s=unicodedata.normalize("NFKD",str(v or ""))
    s="".join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+","_",s.lower()).strip("_")
def finite(v:Any)->float:
    if isinstance(v,bool): raise ValueError(v)
    x=float(v)
    if not math.isfinite(x): raise ValueError(v)
    return x
def sheet_by_headers(wb,required:set[str]):
    for ws in wb.worksheets:
        headers=[norm(c.value) for c in ws[1]]
        if required.issubset(set(headers)): return ws,headers
    raise RuntimeError(f"ISPRA suolo: nessun foglio con {sorted(required)}")
def rows(ws,headers):
    return [{headers[i]:row[i].value for i in range(min(len(headers),len(row)))} for row in ws.iter_rows(min_row=2)]
def main()->None:
    ap=argparse.ArgumentParser(); ap.add_argument("--output",required=True); a=ap.parse_args()
    resp=requests.get(URL,timeout=240,headers={"User-Agent":"OsservatorioVersilia-A3-soil-benchmark/1.0"}); resp.raise_for_status()
    wb=load_workbook(io.BytesIO(resp.content),read_only=True,data_only=True)
    req={"incremento_netto_2023_2024_ettari","suolo_consumato_2024_ettari","suolo_consumato_2024"}
    ws,h=sheet_by_headers(wb,req)
    rr=rows(ws,h)
    name_key=next((x for x in h if "nome_comune" in x or x=="comune"),None)
    code_key=next((x for x in h if "codice_comune" in x or "codice_istat" in x or x=="cod_com"),None)
    if not name_key and not code_key: raise RuntimeError(f"ISPRA suolo: chiave Comune assente {h[:20]}")
    site=json.loads(SITE.read_text(encoding="utf-8"))
    mids=("landUse","landUseChange")
    public={}
    for mid in mids:
        metric=(site.get("metrics") or {}).get(mid) or {}
        items=metric.get("rows") or []
        if len(items)!=7: raise RuntimeError(f"{mid}: pubblico {len(items)}/7")
        public[mid]=items
    by_name={norm(r.get(name_key)):r for r in rr if name_key and r.get(name_key)}
    by_code={re.sub(r"\D","",str(r.get(code_key) or "")):r for r in rr if code_key and r.get(code_key)}
    errors=[]
    for mid,field,tol in (
      ("landUse","suolo_consumato_2024",.11),
      ("landUseChange","incremento_netto_2023_2024_ettari",.011),
    ):
        for row in public[mid]:
            code=re.sub(r"\D","",str(row.get("code") or ""))
            town=norm(row.get("town"))
            src=by_code.get(code) or by_name.get(town)
            if not src:
                errors.append(f"{mid}/{row.get('town')}: Comune ISPRA assente"); continue
            observed=finite(row.get("value")); expected=finite(src.get(field))
            if not math.isclose(observed,expected,rel_tol=0.0,abs_tol=tol):
                errors.append(f"{mid}/{row.get('town')}: {observed} != {expected}")
    if "Regioni_2006_2024" not in wb.sheetnames:
        raise RuntimeError(f"ISPRA suolo: foglio regionale assente {wb.sheetnames}")
    rws=wb["Regioni_2006_2024"]
    rh=[norm(c.value) for c in rws[1]]
    regional=rows(rws,rh)
    reg_name=next((x for x in rh if x=="nome_regione" or x=="regione"),None)
    if not reg_name: raise RuntimeError("ISPRA suolo: colonna Regione assente")
    by_reg={norm(r.get(reg_name)):r for r in regional if r.get(reg_name)}
    if "toscana" not in by_reg:
        raise RuntimeError(f"ISPRA suolo: Toscana assente {list(by_reg)[:30]}")
    tus=by_reg["toscana"]
    if "italia" in by_reg:
        ita=by_reg["italia"]
        italy_land_use=finite(ita["suolo_consumato_2024"])
        italy_change=finite(ita["incremento_netto_2023_2024_ettari"])
        italy_mode="official-row"
    else:
        region_rows=[
            r for key,r in by_reg.items()
            if key not in {"italia","nord","centro","mezzogiorno","sud","isole"}
            and not key.startswith("italia_")
        ]
        if len(region_rows)!=20:
            raise RuntimeError(f"ISPRA suolo: aggregazione Italia non governata, regioni={len(region_rows)} keys={list(by_reg)}")
        consumed=0.0; area=0.0; italy_change=0.0
        for r in region_rows:
            c=finite(r["suolo_consumato_2024_ettari"])
            p=finite(r["suolo_consumato_2024"])
            ch=finite(r["incremento_netto_2023_2024_ettari"])
            if c<0 or not 0<p<=100: raise RuntimeError(f"ISPRA suolo: riga regionale non valida {r.get(reg_name)}")
            consumed+=c
            area+=c*100.0/p
            italy_change+=ch
        if area<=0: raise RuntimeError("ISPRA suolo: superficie nazionale ricostruita nulla")
        italy_land_use=consumed/area*100.0
        italy_mode="20-region-weighted"
    site_metrics=site["metrics"]
    benchmarks={
      "landUse":{"year":"2024","unit":str((site_metrics["landUse"].get("meta") or {}).get("unit") or "percent"),
        "formula":"suolo consumato / superficie territoriale × 100",
        "tuscany":finite(tus["suolo_consumato_2024"]),"italy":italy_land_use},
      "landUseChange":{"year":"2024","unit":str((site_metrics["landUseChange"].get("meta") or {}).get("unit") or "hectares"),
        "formula":"incremento netto 2023-2024 [ettari]",
        "tuscany":finite(tus["incremento_netto_2023_2024_ettari"]),"italy":italy_change},
    }
    gate="PASS" if not errors else "FAIL"
    payload={"schemaVersion":1,"publisher":"ISPRA","profileId":"ispra-consumo-suolo-2024",
      "status":"ACQUIRED_CANDIDATE" if gate=="PASS" else "CANDIDATE_REJECTED","sourceUrl":LANDING,"dataUrl":URL,
      "benchmarks":benchmarks if gate=="PASS" else {},"qualityGate":{"status":gate,"publicReconciliation":"2 metrics × 7/7 PASS" if gate=="PASS" else "FAIL","municipalSheet":ws.title,"regionalSheet":rws.title,"italyAggregation":italy_mode,"errors":errors}}
    out=Path(a.output); out.parent.mkdir(parents=True,exist_ok=True); out.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"status":payload["status"],"candidateMetrics":sorted(payload["benchmarks"]),"gate":payload["qualityGate"]},ensure_ascii=False))
if __name__=="__main__": main()
