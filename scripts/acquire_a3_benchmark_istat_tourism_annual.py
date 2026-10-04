#!/usr/bin/env python3
"""Acquire the native Istat 2024 capacity panel, excluding aggregate rows."""
from __future__ import annotations
import argparse, hashlib, io, json, math, re, urllib.request, zipfile
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
METRIC='tourismBeds'
METRICS=(METRIC,'tourismBedsPer1000','tourismStructuresPer1000')
DEMO=ROOT/'data/source-snapshots/a3-istat-demography-benchmark-2026.json'
URL='https://esploradati.istat.it/databrowser/DWL/Servizi/DCSC%20Capacity%20of%20tourist%20accommodation%20municipal.zip'
STRUCTURE='https://esploradati.istat.it/SDMXWS/rest/dataflow/IT1/DF_BULK_DCSC_CAPACOFTUR/1.0?references=all'
LANDING='https://www.istat.it/informazioni-sulla-rilevazione/capacita-degli-esercizi-ricettivi/'
REGIONS={f'{i:02}0' for i in range(1,21)}-{'040'}|{'041','042'}

def numeric(value):
    if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value) or value<0 or value!=int(value): raise RuntimeError('Capacity missing/noninteger observation')
    return int(value)

def apply_capacity_history(metric, snapshot):
    """Publish only native municipal beds; never infer annual population ratios."""
    history = snapshot.get('municipalBedsHistory')
    if history is None:
        return
    years = history.get('years')
    values = history.get('valuesByCode')
    if years != list(range(2002, 2025)) or not isinstance(values, dict):
        raise RuntimeError('Capacity history period/schema mismatch')
    rows = metric['rows']
    if metric['meta']['key'] != METRIC or {r['code'] for r in rows} != set(values):
        raise RuntimeError('Capacity history municipal scope mismatch')
    if history.get('archiveSha256') != snapshot['source']['archiveSha256']:
        raise RuntimeError('Capacity history archive provenance mismatch')
    for row in rows:
        observations = values[row['code']]
        if not isinstance(observations, list) or len(observations) != len(years):
            raise RuntimeError('Capacity history incomplete municipality')
        for value in observations:
            numeric(value)
        if observations[-1] != row['value']:
            raise RuntimeError('Capacity history latest public value mismatch')
    aggregate = metric['aggregate']
    totals = [sum(values[r['code']][i] for r in rows) for i in range(len(years))]
    if aggregate.get('label') != 'Totale Versilia' or totals[-1] != aggregate['value']:
        raise RuntimeError('Capacity history public aggregate mismatch')
    metadata = {'source': snapshot['publisher'], 'sourceUrl': snapshot['source']['archiveUrl'],
                'sourceSnapshot': 'data/source-snapshots/a3-istat-tourism-capacity-benchmark-2024.json',
                'note': 'Posti letto alberghieri ed extralberghieri nella rilevazione Istat 2002–2024. Il 2025 è escluso per ampliamento del perimetro agli alloggi privati non imprenditoriali.'}
    for holder, observations in [(r, values[r['code']]) for r in rows] + [(aggregate, totals)]:
        candidate = {'years': years, 'values': observations, **metadata}
        existing = holder.get('series')
        if existing not in (None, {}) and existing != candidate:
            raise RuntimeError('Capacity history conflicts with existing public series')
        holder['series'] = candidate

def fetch():
    request=urllib.request.Request(URL,headers={'User-Agent':'OsservatorioVersilia-A3-capacity/1.0'})
    with urllib.request.urlopen(request,timeout=300) as response:
        body=response.read(80*1024*1024+1)
    if len(body)>80*1024*1024: raise RuntimeError('Capacity archive exceeds 80 MiB')
    return body

def parse_capacity_history(body, codes):
    """Read the seven native time series, with no territorial reconstruction."""
    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        workbook = load_workbook(io.BytesIO(archive.read('Capacità comunale 2002-2025.xlsx')),
                                 read_only=True, data_only=True)
    years = list(range(2002, 2025))
    values = {code: {} for code in codes}
    try:
        for row in workbook.active.iter_rows(min_row=7, max_col=56, values_only=True):
            year, code = row[0], str(row[7])
            if not isinstance(year, int) or year not in years or code not in values:
                continue
            if year in values[code]:
                raise RuntimeError('Capacity history duplicate municipal observation')
            hotel, other, total = [numeric(row[i]) for i in (33, 53, 55)]
            if hotel + other != total:
                raise RuntimeError('Capacity history component total mismatch')
            values[code][year] = total
    finally:
        workbook.close()
    if any(set(observations) != set(years) for observations in values.values()):
        raise RuntimeError('Capacity history incomplete native series')
    return {'years': years,
            'valuesByCode': {code: [values[code][year] for year in years] for code in sorted(codes)},
            'archiveSha256': hashlib.sha256(body).hexdigest()}

def parse_archive(body):
    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(body)) as z:
        names=[n for n in z.namelist() if n.endswith('.xlsx')]
        if names!=['Capacità comunale 2002-2025.xlsx']: raise RuntimeError('Capacity archive member mismatch')
        xlsx=z.read(names[0])
    workbook=load_workbook(io.BytesIO(xlsx),read_only=True,data_only=True)
    if workbook.sheetnames!=['Capacità ricettiva']: raise RuntimeError('Capacity worksheet mismatch')
    rows=workbook.active.iter_rows(max_col=56,values_only=True)
    headers=[next(rows) for _ in range(6)]
    if 'TOTALE' not in str(headers[2][54]) or 'totale alberghi' not in str(headers[4][32]) or 'totale extra-alberghieri' not in str(headers[4][52]): raise RuntimeError('Capacity total/component column mismatch')
    if headers[5][0]!='Anno/Year' or headers[5][7]!='Cod. Istat' or any('Beds' not in str(headers[5][i]) for i in (33,53,55)): raise RuntimeError('Capacity year/code/unit mismatch')
    records=[]; provinces=[]; national=[]; previous=2026
    for row in rows:
        year=row[0]
        if not isinstance(year,int): continue
        if year>previous: raise RuntimeError('Capacity year order changed')
        previous=year
        if year<2024: break
        if year!=2024: continue
        hotel,other,total=[numeric(row[i]) for i in (33,53,55)]
        structures=numeric(row[54])
        if hotel+other!=total: raise RuntimeError('Capacity hotel/other total mismatch')
        code,region,province=row[7],row[2],row[4]
        if code is not None:
            records.append([str(code),str(region),hotel,other,total,structures])
        elif region is not None:
            if row[5]!='TOTALE' or not isinstance(province,str) or not re.fullmatch(r'\d{3}',province): raise RuntimeError('Capacity province aggregate mismatch')
            provinces.append([province,str(region),hotel,other,total,structures])
        else:
            national.append({'hotelBeds':hotel,'otherBeds':other,'beds':total,'structures':structures})
    workbook.close()
    if len(national)!=1 or len(provinces)!=107: raise RuntimeError('Capacity aggregate scope mismatch')
    return sorted(records),sorted(provinces),national[0],{'archiveMember':names[0],'archiveSha256':hashlib.sha256(body).hexdigest(),'workbookSha256':hashlib.sha256(xlsx).hexdigest(),'dataflowStructureUrl':STRUCTURE}

def verify_panel(records,provinces,national):
    if len(records)!=7899 or len(provinces)!=107: raise RuntimeError('Capacity incomplete native panel')
    seen=set();regions=set();by_province={}
    for code,region,hotel,other,total,structures in records:
        if not isinstance(code,str) or not re.fullmatch(r'\d{6}',code) or code in seen or code.endswith('000') or region not in REGIONS: raise RuntimeError('Capacity invalid/duplicate municipality')
        for value in (hotel,other,total,structures):numeric(value)
        if hotel+other!=total:raise RuntimeError('Capacity record components mismatch')
        seen.add(code);regions.add(region)
        p=by_province.setdefault(code[:3],[0,0,0,0]);p[0]+=hotel;p[1]+=other;p[2]+=total;p[3]+=structures
    if regions!=REGIONS or sum(r[1]=='090' for r in records)!=273:raise RuntimeError('Capacity region/Tuscany scope mismatch')
    seen_prov=set()
    for province,region,hotel,other,total,structures in provinces:
        for value in (hotel,other,total,structures): numeric(value)
        if province in seen_prov or region not in REGIONS or by_province.get(province)!=[hotel,other,total,structures]:raise RuntimeError('Capacity native province reconciliation mismatch')
        if any(r[1]!=region for r in records if r[0][:3]==province):raise RuntimeError('Capacity province-region mapping mismatch')
        seen_prov.add(province)
    if seen_prov!=set(by_province):raise RuntimeError('Capacity province coverage mismatch')
    expected={'hotelBeds':sum(r[2] for r in records),'otherBeds':sum(r[3] for r in records),'beds':sum(r[4] for r in records),'structures':sum(r[5] for r in records)}
    for value in national.values(): numeric(value)
    if national!=expected or national['beds']!=5498773:raise RuntimeError('Capacity official national total mismatch')

def validate_snapshot(metric,snapshot,population=None):
    mid=metric['meta']['key'];unit='number' if mid==METRIC else 'per1000'
    if mid not in METRICS:raise RuntimeError('Capacity metric mismatch')
    if snapshot.get('qualityGate')!={'status':'PASS','errors':[],'publicReconciliation':'3 metrics × 7/7 PASS'}:raise RuntimeError('Capacity gate not PASS')
    if metric['meta'].get('year')!='2024' or metric['meta'].get('unit')!=unit:raise RuntimeError('Capacity public year/unit mismatch')
    if snapshot.get('referenceYear')!=2024 or snapshot.get('sourceUrl')!=LANDING or snapshot['source']['archiveUrl']!=URL:raise RuntimeError('Capacity source/reference mismatch')
    records=snapshot['records'];verify_panel(records,snapshot['provinceTotals'],snapshot['officialNational'])
    demo_bytes=DEMO.read_bytes();demo=json.loads(demo_bytes)['benchmarks']['population']
    if snapshot['population']['sha256']!=hashlib.sha256(demo_bytes).hexdigest() or snapshot['population']['year']!='2026' or demo['year']!='2026':raise RuntimeError('Capacity population provenance mismatch')
    for target in METRICS:
        benchmark=snapshot['benchmarks'][target];target_unit='number' if target==METRIC else 'per1000'
        if benchmark.get('year')!='2024' or benchmark.get('unit')!=target_unit:raise RuntimeError('Capacity benchmark year/unit mismatch')
        column=5 if target=='tourismStructuresPer1000' else 4
        for scope in ('tuscany','italy'):
            total=sum(r[column] for r in records if scope=='italy' or r[1]=='090')
            expected=total if target==METRIC else total/demo[scope]*1000
            value=benchmark[scope]
            if isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isclose(value,expected,abs_tol=1e-10,rel_tol=0):raise RuntimeError('Capacity benchmark aggregate mismatch')
    source={r[0]:r for r in records};rows=metric['rows'];proof=snapshot['municipalReconciliation']
    if len(rows)!=7 or len({r['code'] for r in rows})!=7 or {r['code'] for r in rows}!=set(proof):raise RuntimeError('Capacity public scope not 7/7')
    for row in rows:
        record=source[row['code']]
        if proof[row['code']]!={'hotelBeds':record[2],'otherBeds':record[3],'beds':record[4],'structures':record[5]} or record[1]!='090':raise RuntimeError('Capacity 7/7 component mismatch')
        if mid==METRIC:expected=record[4]
        else:
            if population is None or population['meta']['year']!='2026':raise RuntimeError('Capacity public population year mismatch')
            pop_rows=[r for r in population['rows'] if r['code']==row['code']]
            if len(pop_rows)!=1 or pop_rows[0]['value']<=0 or isinstance(pop_rows[0]['value'],bool):raise RuntimeError('Capacity public population missing')
            expected=record[5 if mid=='tourismStructuresPer1000' else 4]/pop_rows[0]['value']*1000
        if isinstance(row['value'],bool) or not math.isclose(row['value'],expected,rel_tol=0,abs_tol=1e-9):raise RuntimeError('Capacity public value mismatch')

def main():
    p=argparse.ArgumentParser();p.add_argument('--input-zip',type=Path);p.add_argument('--output',type=Path,required=True);p.add_argument('--include-history',action='store_true');args=p.parse_args()
    body=args.input_zip.read_bytes() if args.input_zip else fetch();records,provinces,national,source=parse_archive(body)
    site=json.loads((ROOT/'data/site-data.json').read_text());metric=site['metrics'][METRIC];source['archiveUrl']=URL
    by_code={r[0]:r for r in records};proof={r['code']:{'hotelBeds':by_code[r['code']][2],'otherBeds':by_code[r['code']][3],'beds':by_code[r['code']][4],'structures':by_code[r['code']][5]} for r in metric['rows']}
    demo_bytes=DEMO.read_bytes();demo=json.loads(demo_bytes)['benchmarks']['population'];benchmarks={}
    for target in METRICS:
        column=5 if target=='tourismStructuresPer1000' else 4
        values={scope:sum(r[column] for r in records if scope=='italy' or r[1]=='090') for scope in ('tuscany','italy')}
        benchmarks[target]={'year':'2024','unit':'number' if target==METRIC else 'per1000',**{scope:total if target==METRIC else total/demo[scope]*1000 for scope,total in values.items()}}
    snapshot={'schemaVersion':1,'profileId':'istat-tourism-annual','referenceYear':2024,'publisher':'Istat — Capacità degli esercizi ricettivi','sourceUrl':LANDING,'source':source,'records':records,'provinceTotals':provinces,'officialNational':national,'municipalReconciliation':proof,'benchmarks':benchmarks,'population':{'sourceSnapshot':str(DEMO.relative_to(ROOT)),'sha256':hashlib.sha256(demo_bytes).hexdigest(),'year':'2026'},
        'scope':{'note':'Posti letto alberghieri ed extralberghieri 2024 nel perimetro nativo Istat: 7.899 record comunali, 273 in Toscana. Totali provinciali e nazionale esclusi dalla somma comunale e usati soltanto per riconciliazione (107/107 province e totale Italia). Il perimetro conserva le ripartizioni comunali pubblicate nella rilevazione, senza conversione alla geografia delle dichiarazioni MEF. Dal 2025 la rilevazione amplia il perimetro agli alloggi privati non imprenditoriali: non vengono sostituiti i dati 2024 con quelli 2025.'},
        'qualityGate':{'status':'PASS','errors':[],'publicReconciliation':'3 metrics × 7/7 PASS'}}
    snapshot['scope']['note']+=' Per i rapporti ogni 1.000 residenti si conserva e si dichiara la popolazione pubblica al 1° gennaio 2026, applicando lo stesso denominatore temporale anche a Toscana/Italia.'
    for target in METRICS:validate_snapshot(site['metrics'][target],snapshot,site['metrics']['population'])
    if args.include_history:
        snapshot['municipalBedsHistory'] = parse_capacity_history(body, set(proof))
        apply_capacity_history(metric, snapshot)
    args.output.parent.mkdir(parents=True,exist_ok=True);args.output.write_text(json.dumps(snapshot,ensure_ascii=False,separators=(',',':'))+'\n')
    print(f'Capacity candidates PASS: {benchmarks}; 7899 records; 107 provinces and 3 metrics × 7/7 public. Not ACQUIRED before publication gate.')

if __name__=='__main__':main()
