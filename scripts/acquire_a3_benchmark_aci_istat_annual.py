#!/usr/bin/env python3
from __future__ import annotations

import argparse,io,json,math
from pathlib import Path
from typing import Any
import requests
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
SITE=ROOT/"data/site-data.json"
URL="https://www.istat.it/wp-content/uploads/2026/02/PARCO_VEICOLARE_2024.xlsx"
LANDING="https://www.istat.it/comunicato-stampa/indicatori-del-parco-veicolare-anno-2024/"

def finite(v:Any)->float:
    if isinstance(v,bool): raise ValueError(v)
    x=float(v)
    if not math.isfinite(x): raise ValueError(v)
    return x

def find_geo(ws,name:str)->tuple[Any,...]:
    target=name.casefold()
    hits=[]
    for row in ws.iter_rows(values_only=True):
        label=str(row[2] or "").strip().casefold() if len(row)>2 else ""
        if label==target or (target=="italia" and label.startswith("italia")):
            hits.append(row)
    if len(hits)!=1:
        raise RuntimeError(f"ACI/Istat {ws.title}: riga {name} non univoca ({len(hits)})")
    return hits[0]

def reconcile_public_rows(ws,rows:list[dict],value_index:int,metric_id:str,tolerance:float)->list[dict]:
    evidence=[]
    errors=[]
    for item in rows:
        town=str(item.get("town") or "").strip()
        if not town:
            errors.append(f"{metric_id}: town pubblico assente")
            continue
        try:
            source=find_geo(ws,town)
            expected=finite(source[value_index])
            observed=finite(item.get("value"))
        except Exception as exc:
            errors.append(f"{metric_id}/{town}: {type(exc).__name__}: {exc}")
            continue
        evidence.append({"town":town,"observed":observed,"source":expected})
        if not math.isclose(observed,expected,rel_tol=0.0,abs_tol=tolerance):
            errors.append(f"{metric_id}/{town}: {observed} != {expected}")
    if errors:
        raise RuntimeError("; ".join(errors[:20]))
    if len(evidence)!=7:
        raise RuntimeError(f"{metric_id}: riconciliazione comunale {len(evidence)}/7")
    return evidence

def public_contract(site:dict,metric_id:str,unit:str,year:str,formula_tokens:tuple[str,...])->list[dict]:
    metric=(site.get("metrics") or {}).get(metric_id) or {}
    meta=metric.get("meta") or {}
    rows=metric.get("rows") or []
    if len(rows)!=7 or any(not isinstance(r.get("value"),(int,float)) for r in rows):
        raise RuntimeError(f"{metric_id}: pubblico non numerico 7/7")
    if str(meta.get("unit") or "")!=unit or str(meta.get("year") or "")!=year:
        raise RuntimeError(f"{metric_id}: contratto unit/year inatteso {meta.get('unit')}/{meta.get('year')}")
    source=str(meta.get("source") or "").casefold()
    if "aci" not in source or "istat" not in source:
        raise RuntimeError(f"{metric_id}: fonte pubblica inattesa {meta.get('source')}")
    formula=str((metric.get("method") or {}).get("formula") or "").casefold()
    if not all(token.casefold() in formula for token in formula_tokens):
        raise RuntimeError(f"{metric_id}: formula pubblica inattesa {formula}")
    return rows

def main()->None:
    ap=argparse.ArgumentParser(); ap.add_argument("--output",required=True); a=ap.parse_args()
    site=json.loads(SITE.read_text(encoding="utf-8"))
    motor_rows=public_contract(site,"motorization","per1000","2024",("autovetture","popolazione","1.000"))
    polluting_rows=public_contract(site,"pollutingCars","percent","2024",("euro 0","autovetture totali","100"))
    if {str(r.get("code") or "") for r in motor_rows}!={str(r.get("code") or "") for r in polluting_rows}:
        raise RuntimeError("ACI/Istat: perimetri comunali pubblici incoerenti")

    response=requests.get(URL,timeout=240,headers={"User-Agent":"OsservatorioVersilia-A3-ACI-benchmark/1.0"})
    response.raise_for_status()
    wb=load_workbook(io.BytesIO(response.content),read_only=True,data_only=True)
    if "mobilità urbana_11.1" not in wb.sheetnames or "mobilità urbana_12.1" not in wb.sheetnames:
        raise RuntimeError(f"ACI/Istat: fogli attesi assenti {wb.sheetnames}")
    ws_motor=wb["mobilità urbana_11.1"]; ws_poll=wb["mobilità urbana_12.1"]
    motor_municipal=reconcile_public_rows(ws_motor,motor_rows,5,"motorization",.11)
    poll_municipal=reconcile_public_rows(ws_poll,polluting_rows,19,"pollutingCars",.11)
    m_tus=find_geo(ws_motor,"Toscana"); m_ita=find_geo(ws_motor,"Italia")
    p_tus=find_geo(ws_poll,"Toscana"); p_ita=find_geo(ws_poll,"Italia")
    motor_tus=finite(m_tus[5]); motor_ita=finite(m_ita[5])
    poll_tus=finite(p_tus[19]); poll_ita=finite(p_ita[19])
    if not (300<=motor_tus<=1200 and 300<=motor_ita<=1200):
        raise RuntimeError("ACI/Istat: tasso motorizzazione fuori intervallo")
    if not (0<=poll_tus<=100 and 0<=poll_ita<=100):
        raise RuntimeError("ACI/Istat: quota Euro 0-3 fuori intervallo")

    payload={
      "schemaVersion":1,
      "publisher":"ACI / Istat — parco veicolare",
      "profileId":"aci-istat-annual",
      "status":"ACQUIRED_CANDIDATE",
      "sourceUrl":LANDING,"dataUrl":URL,
      "benchmarks":{
        "motorization":{"year":"2024","unit":"per1000","formula":"autovetture / popolazione residente × 1.000","tuscany":motor_tus,"italy":motor_ita},
        "pollutingCars":{"year":"2024","unit":"percent","formula":"autovetture Euro 0–3 / autovetture totali × 100","tuscany":poll_tus,"italy":poll_ita},
      },
      "qualityGate":{
        "status":"PASS",
        "publicContractCompatibility":"2 metrics × 7/7 public rows; source/unit/year/formula PASS",
        "publicReconciliation":"2 metrics × 7/7 municipal rows PASS",
        "officialBenchmarkRows":"Toscana + Italia 2024 from Istat tables 11.1 and 12.1 PASS",
        "errors":[]
      },
      "evidence":{
        "motorization":{"sheet":"mobilità urbana_11.1","municipalRows":motor_municipal,"tuscanyRow":list(m_tus),"italyRow":list(m_ita)},
        "pollutingCars":{"sheet":"mobilità urbana_12.1","municipalRows":poll_municipal,"tuscanyRow":list(p_tus),"italyRow":list(p_ita)}
      }
    }
    out=Path(a.output); out.parent.mkdir(parents=True,exist_ok=True)
    out.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"status":payload["status"],"candidateMetrics":sorted(payload["benchmarks"])},ensure_ascii=False))
if __name__=="__main__": main()
