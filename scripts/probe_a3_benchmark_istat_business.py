#!/usr/bin/env python3
from __future__ import annotations

import argparse
import io
import json
import zipfile
from pathlib import Path

import requests
from openpyxl import load_workbook

SOURCE_URL = "https://www.istat.it/wp-content/uploads/2026/01/Tavole.zip"


def row_values(sheet, row_index: int, max_col: int = 40) -> list:
    return [sheet.cell(row=row_index, column=col).value for col in range(1, max_col + 1)]


def compact(values: list) -> list:
    last = 0
    for idx, value in enumerate(values, start=1):
        if value not in (None, ""):
            last = idx
    return values[:last]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    session = requests.Session()
    session.headers["User-Agent"] = "OsservatorioVersilia-A3-benchmark-probe/1.0"
    response = session.get(SOURCE_URL, timeout=240)
    response.raise_for_status()
    body = response.content

    report = {
        "schemaVersion": 1,
        "sourceUrl": SOURCE_URL,
        "archiveBytes": len(body),
        "members": [],
        "candidateSheets": [],
    }

    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        report["members"] = archive.namelist()
        workbook_members = [
            name for name in archive.namelist()
            if name.lower().endswith((".xlsx", ".xlsm"))
        ]
        for member in workbook_members:
            raw = archive.read(member)
            try:
                workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            except Exception as exc:
                report["candidateSheets"].append({
                    "member": member,
                    "loadError": f"{type(exc).__name__}: {exc}",
                })
                continue
            for sheet in workbook.worksheets:
                matches = []
                max_row = min(sheet.max_row or 0, 220)
                max_col = min(sheet.max_column or 0, 50)
                for row_index in range(1, max_row + 1):
                    values = row_values(sheet, row_index, max_col)
                    text = " | ".join(str(value) for value in values if value not in (None, ""))
                    lowered = text.lower()
                    if "toscana" in lowered or lowered.strip() == "italia" or " italia " in f" {lowered} ":
                        matches.append({
                            "row": row_index,
                            "values": compact(values),
                        })
                if matches:
                    report["candidateSheets"].append({
                        "member": member,
                        "sheet": sheet.title,
                        "maxRow": sheet.max_row,
                        "maxColumn": sheet.max_column,
                        "headerRows": [compact(row_values(sheet, header_row, max_col)) for header_row in range(1, min(12, max_row) + 1)],\n                            "matches": matches[:20],
                    })

    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "archiveBytes": report["archiveBytes"],
        "memberCount": len(report["members"]),
        "candidateSheetCount": len(report["candidateSheets"]),
        "output": str(output),
    }, ensure_ascii=False, indent=2))

    if not report["candidateSheets"]:
        raise SystemExit("Nessuna tavola ISTAT con Toscana/Italia individuata nel pacchetto ufficiale")


if __name__ == "__main__":
    main()
