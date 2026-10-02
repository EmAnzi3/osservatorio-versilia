#!/usr/bin/env python3
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import zipfile
from html import unescape
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlparse

import requests
from openpyxl import load_workbook

FILE_RE=re.compile(r"""(?:href|src)=["']([^"']+\.(?:csv|json|geojson|zip|xlsx?|ods)(?:\?[^"']*)?)["']""",re.I)
ABS_RE=re.compile(r"""https?://[^\s"'<>]+\.(?:csv|json|geojson|zip|xlsx?|ods)(?:\?[^\s"'<>]*)?""",re.I)
YEAR_RE=re.compile(r"\b(?:19|20)\d{2}(?:\s*[-/]\s*(?:19|20)?\d{2})?\b")
TARGET_WORDS=("toscana","italia","italy","nazionale","national")
MAX_BODY=60*1024*1024


def load(path:Path)->dict[str,Any]:
    value=json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value,dict):
        raise RuntimeError(f"JSON non-oggetto: {path}")
    return value


def profile_bundle(manifest:dict[str,Any],profile_id:str)->dict[str,Any]:
    for item in manifest.get("profiles") or []:
        if isinstance(item,dict) and item.get("profileId")==profile_id:
            return item
    raise RuntimeError(f"Profilo {profile_id} assente dal manifest fan-out")


def registry_urls(registry:dict[str,Any],profile_id:str)->list[str]:
    urls=[]
    for mapping_name in ("sourceProfileByUrl","sourceUrlProfiles"):
        mapping=registry.get(mapping_name)
        if not isinstance(mapping,dict):
            continue
        for url,value in mapping.items():
            matched=value==profile_id or (isinstance(value,list) and profile_id in value)
            if matched and str(url).startswith("http"):
                urls.append(str(url))
    return urls


def safe_get(session:requests.Session,url:str)->requests.Response:
    response=session.get(url,timeout=120,allow_redirects=True,stream=True)
    response.raise_for_status()
    chunks=[]; total=0
    for chunk in response.iter_content(1024*128):
        if not chunk: continue
        total+=len(chunk)
        if total>MAX_BODY:
            response.close()
            raise RuntimeError(f"download oltre limite {MAX_BODY} byte")
        chunks.append(chunk)
    response._content=b"".join(chunks)
    response._content_consumed=True
    return response


def text_from(response:requests.Response)->str:
    try:
        return response.content.decode(response.encoding or "utf-8",errors="replace")
    except Exception:
        return response.content.decode("utf-8",errors="replace")


def csv_probe(body:bytes)->dict[str,Any]:
    text=body.decode("utf-8-sig",errors="replace")
    sample=text[:8192]
    try:
        delimiter=csv.Sniffer().sniff(sample,delimiters=",;|\t").delimiter
    except csv.Error:
        delimiter=","
    reader=csv.reader(io.StringIO(text),delimiter=delimiter)
    rows=[]
    for index,row in enumerate(reader):
        rows.append(row[:80])
        if index>=5: break
    return {"delimiter":delimiter,"sampleRows":rows}


def xlsx_probe(body:bytes)->dict[str,Any]:
    workbook=load_workbook(io.BytesIO(body),read_only=True,data_only=True)
    sheets=[]
    for sheet in workbook.worksheets[:25]:
        sample=[]; matches=[]
        for r_index,row in enumerate(sheet.iter_rows(values_only=True),start=1):
            values=[value for value in list(row)[:80]]
            if r_index<=8:
                sample.append(values)
            joined=" | ".join(str(v) for v in values if v not in (None,"")).lower()
            if any(word in joined for word in TARGET_WORDS):
                matches.append({"row":r_index,"values":values})
                if len(matches)>=12: break
            if r_index>=400 and not matches:
                break
        sheets.append({
            "title":sheet.title,
            "maxRow":sheet.max_row,
            "maxColumn":sheet.max_column,
            "sampleRows":sample,
            "geoMatches":matches,
        })
    return {"sheets":sheets}


def file_priority(profile_id:str,url:str)->tuple[int,str]:
    text=url.lower()
    score=0
    for year,points in (("2026",12),("2025",11),("2024",10),("2023",9),("2022",6),("2021",5)):
        if year in text:
            score+=points
            break
    if "toscana" in text or "/reg_" in text or "dati_regionali" in text:
        score+=25
    if "italia" in text or "national" in text or "nazionale" in text:
        score+=18
    if text.endswith(".xlsx") or ".xlsx?" in text:
        score+=7
    elif text.endswith(".zip") or ".zip?" in text:
        score+=6
    elif text.endswith(".csv") or ".csv?" in text:
        score+=5
    elif text.endswith(".json") or ".json?" in text:
        score+=3

    if profile_id=="istat-census-annual":
        if "dati_regionali_2023" in text: score+=80
        if "comuni_2023" in text: score+=45
    elif profile_id=="istat-demography-annual":
        if "046_lucca" in text: score+=80
        if "comuni" in text: score+=55
        if "/p2_" in text or "/p02" in text: score+=20
        if "posas" in text: score+=15
    elif profile_id=="mim-school-year":
        if "202425" in text: score+=70
        if any(token in text for token in ("alucorsoindcla","alutemposcuola","scuanagrafe")): score+=30
    elif profile_id=="mef-irpef-annual":
        if "/reg_" in text: score+=70
        if "base_comunale_csv_2024" in text: score+=65
        if "tipo_reddito_2024" in text or "calcolo_irpef_2024" in text: score+=35
    elif profile_id=="openbdap-annual":
        score+=50
    elif profile_id=="istat-business-annual":
        if "tavole.zip" in text: score+=80
    elif profile_id=="istat-agriculture-census-2020":
        if "manifest.json" in text: score+=50
    return (-score,url)


def csv_geo_probe(body:bytes,max_rows:int=12000)->dict[str,Any]:
    text=body.decode("utf-8-sig",errors="replace")
    sample=text[:8192]
    try:
        delimiter=csv.Sniffer().sniff(sample,delimiters=",;|\t").delimiter
    except csv.Error:
        delimiter=","
    reader=csv.reader(io.StringIO(text),delimiter=delimiter)
    sample_rows=[]; geo_matches=[]; rows_scanned=0
    for index,row in enumerate(reader):
        rows_scanned=index+1
        values=row[:100]
        if index<6:
            sample_rows.append(values)
        joined=" | ".join(str(v) for v in values if v not in (None,"")).lower()
        if any(word in joined for word in TARGET_WORDS):
            geo_matches.append({"row":index+1,"values":values})
            if len(geo_matches)>=25 and index>=100:
                break
        if index+1>=max_rows:
            break
    return {
        "delimiter":delimiter,
        "sampleRows":sample_rows,
        "geoMatches":geo_matches,
        "rowsScanned":rows_scanned,
    }


def zip_probe(body:bytes)->dict[str,Any]:
    with zipfile.ZipFile(io.BytesIO(body)) as archive:
        names=archive.namelist()
        structured=[
            name for name in names
            if name.lower().endswith((".csv",".json",".geojson",".xlsx",".xlsm",".ods"))
        ]
        inspected=[]
        def member_score(name:str)->tuple[int,str]:
            low=name.lower()
            score=0
            if any(token in low for token in ("toscana","region","italia","national","nazionale")): score+=30
            if "2025" in low: score+=12
            elif "2024" in low: score+=11
            elif "2023" in low: score+=10
            if low.endswith(".csv"): score+=5
            elif low.endswith((".xlsx",".xlsm")): score+=4
            elif low.endswith((".json",".geojson")): score+=3
            return (-score,name)
        for name in sorted(structured,key=member_score)[:10]:
            try:
                raw=archive.read(name)
                low=name.lower()
                if low.endswith(".csv"):
                    detail={"member":name,"bytes":len(raw),"kind":"csv",**csv_geo_probe(raw)}
                elif low.endswith((".xlsx",".xlsm")):
                    detail={"member":name,"bytes":len(raw),"kind":"xlsx",**xlsx_probe(raw)}
                elif low.endswith((".json",".geojson")):
                    value=json.loads(raw.decode("utf-8-sig",errors="replace"))
                    detail={"member":name,"bytes":len(raw),"kind":"json"}
                    if isinstance(value,dict):
                        detail["keys"]=list(value)[:80]
                    elif isinstance(value,list):
                        detail["length"]=len(value); detail["sample"]=value[:2]
                else:
                    detail={"member":name,"bytes":len(raw),"kind":"other"}
            except Exception as exc:
                detail={"member":name,"kind":"parse-error","error":f"{type(exc).__name__}: {exc}"}
            inspected.append(detail)
        return {
            "memberCount":len(names),
            "members":names[:120],
            "structuredMembers":structured[:120],
            "inspectedMembers":inspected,
        }


def structured_probe(url:str,response:requests.Response)->dict[str,Any]:
    content_type=(response.headers.get("content-type") or "").lower()
    path=urlparse(response.url).path.lower()
    body=response.content
    try:
        if path.endswith((".xlsx",".xlsm")) or "spreadsheetml" in content_type:
            return {"kind":"xlsx",**xlsx_probe(body)}
        if path.endswith(".zip") or "zip" in content_type:
            return {"kind":"zip",**zip_probe(body)}
        if path.endswith((".json",".geojson")) or "json" in content_type:
            value=json.loads(body.decode("utf-8-sig",errors="replace"))
            if isinstance(value,dict):
                return {"kind":"json","topLevel":"object","keys":list(value)[:80]}
            if isinstance(value,list):
                return {"kind":"json","topLevel":"array","length":len(value),"sample":value[:2]}
        if path.endswith(".csv") or "csv" in content_type or "text/plain" in content_type:
            return {"kind":"csv",**csv_geo_probe(body)}
    except Exception as exc:
        return {"kind":"structured-parse-error","error":f"{type(exc).__name__}: {exc}"}
    return {"kind":"other"}


def probe_url(session:requests.Session,url:str)->dict[str,Any]:
    try:
        response=safe_get(session,url)
    except Exception as exc:
        return {"requestedUrl":url,"ok":False,"error":f"{type(exc).__name__}: {exc}"}

    content_type=response.headers.get("content-type") or ""
    result={
        "requestedUrl":url,
        "resolvedUrl":response.url,
        "ok":True,
        "status":response.status_code,
        "contentType":content_type,
        "bytes":len(response.content),
        "etag":response.headers.get("etag"),
        "lastModified":response.headers.get("last-modified"),
    }
    lower_ct=content_type.lower()
    if "html" in lower_ct or "text" in lower_ct:
        text=text_from(response)
        links={urljoin(response.url,unescape(match)) for match in FILE_RE.findall(text)}
        links.update(unescape(match) for match in ABS_RE.findall(text))
        result["yearsMentioned"]=sorted(set(YEAR_RE.findall(text)))[:80]
        result["fileLinks"]=sorted(links)[:120]
        result["textMentions"]={word:(word in text.lower()) for word in TARGET_WORDS}
    result["structured"]=structured_probe(url,response)
    return result


def main()->int:
    ap=argparse.ArgumentParser()
    ap.add_argument("--manifest",type=Path,required=True)
    ap.add_argument("--registry",type=Path,required=True)
    ap.add_argument("--profile",required=True)
    ap.add_argument("--output",type=Path,required=True)
    ap.add_argument("--deep-files",type=int,default=0)
    args=ap.parse_args()

    manifest=load(args.manifest)
    registry=load(args.registry)
    bundle=profile_bundle(manifest,args.profile)
    urls=[]
    urls.extend(
        str(ref) for ref in bundle.get("sourceReferences") or []
        if str(ref).startswith("http")
    )
    urls.extend(registry_urls(registry,args.profile))
    urls=list(dict.fromkeys(urls))

    session=requests.Session()
    session.headers["User-Agent"]="OsservatorioVersilia-A3-parallel-source-worker/1.0"
    probes=[probe_url(session,url) for url in urls]

    file_links=[]
    for probe in probes:
        for link in probe.get("fileLinks") or []:
            if link not in file_links:
                file_links.append(link)

    selected_deep=sorted(file_links,key=lambda url:file_priority(args.profile,url))[:max(0,args.deep_files)]
    deep_probes=[probe_url(session,url) for url in selected_deep]

    status="SOURCE_UNREACHABLE"
    if any(item.get("ok") for item in probes):
        status="LANDING_SOURCE_REACHABLE"
    if any((item.get("structured") or {}).get("kind") in {"csv","json","xlsx","zip"} for item in probes):
        status="STRUCTURED_SOURCE_REACHED"
    if file_links:
        status="STRUCTURED_FILES_DISCOVERED"
    if any(
        item.get("ok") and (item.get("structured") or {}).get("kind") in {"csv","json","xlsx","zip"}
        for item in deep_probes
    ):
        status="DEEP_STRUCTURED_SOURCE_REACHED"

    payload={
        "schemaVersion":1,
        "dimension":manifest.get("dimension"),
        "profileId":args.profile,
        "publisher":bundle.get("publisher"),
        "rank":bundle.get("rank"),
        "pairCount":bundle.get("pairCount"),
        "metricIds":bundle.get("metricIds"),
        "sourceReferences":bundle.get("sourceReferences"),
        "registryUrls":registry_urls(registry,args.profile),
        "status":status,
        "probeCount":len(probes),
        "reachableCount":sum(1 for item in probes if item.get("ok")),
        "discoveredFileCount":len(file_links),
        "discoveredFiles":file_links[:120],
        "deepFileCount":len(selected_deep),
        "deepFiles":selected_deep,
        "deepProbes":deep_probes,
        "probes":probes,
    }
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(payload,ensure_ascii=False,indent=2,default=str)+"\n",encoding="utf-8")
    print(
        f"A3 fan-out worker {args.profile}: {status} · "
        f"{payload['reachableCount']}/{payload['probeCount']} URL raggiungibili · "
        f"{payload['discoveredFileCount']} file candidati · "
        f"{payload['deepFileCount']} deep-probe · "
        f"{payload['pairCount']} pair."
    )
    # Source discovery failures are evidence, not CI infrastructure failures.
    # The artifact must always be produced so fan-in can classify the profile.
    return 0


if __name__=="__main__":
    raise SystemExit(main())
